from __future__ import annotations

import csv
import os
from pathlib import Path
from typing import Any

import geopandas as gpd
from openpyxl import load_workbook

from .input import id_key
from .models import Node, RouteResult, Settings


SUMMARY_FIELDS = [
    "mode",
    "from_id",
    "to_id",
    "from_name",
    "to_name",
    "country_code",
    "status",
    "message",
    "network_length_m",
    "metric_distance_m",
    "from_snap_distance_m",
    "to_snap_distance_m",
]


def _successful_metrics(results: list[RouteResult]) -> dict[tuple[str, str], float]:
    metrics: dict[tuple[str, str], float] = {}
    for result in results:
        if result.status == "success" and result.metric_distance_m is not None:
            metrics[(id_key(result.from_id), id_key(result.to_id))] = (
                result.metric_distance_m / 1000.0
            )
    return metrics


def update_truck_sheet(
    settings: Settings,
    nodes: list[Node],
    results: list[RouteResult],
) -> Path:
    workbook = load_workbook(settings.workbook_path)
    if settings.truck_sheet in workbook.sheetnames:
        index = workbook.sheetnames.index(settings.truck_sheet)
        worksheet = workbook[settings.truck_sheet]
        workbook.remove(worksheet)
        worksheet = workbook.create_sheet(settings.truck_sheet, index)
    else:
        worksheet = workbook.create_sheet(settings.truck_sheet)

    metrics = _successful_metrics(results)
    worksheet.cell(row=1, column=1, value="from_id")
    for column, node in enumerate(nodes, start=2):
        worksheet.cell(row=1, column=column, value=node.node_id)
    for row, origin in enumerate(nodes, start=2):
        worksheet.cell(row=row, column=1, value=origin.node_id)
        for column, destination in enumerate(nodes, start=2):
            value = metrics.get(
                (id_key(origin.node_id), id_key(destination.node_id)),
                0.0,
            )
            cell = worksheet.cell(row=row, column=column, value=value)
            cell.number_format = "0.000"

    destination = settings.output_workbook_path
    destination.parent.mkdir(parents=True, exist_ok=True)
    temporary = destination.with_name(
        f".{destination.stem}.truck-router.tmp{destination.suffix}"
    )
    workbook.save(temporary)
    os.replace(temporary, destination)
    return destination


def write_summary(settings: Settings, results: list[RouteResult]) -> Path:
    settings.output_dir.mkdir(parents=True, exist_ok=True)
    path = settings.output_dir / settings.summary_filename
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=SUMMARY_FIELDS)
        writer.writeheader()
        for result in results:
            writer.writerow(
                {
                    "mode": "truck",
                    "from_id": result.from_id,
                    "to_id": result.to_id,
                    "from_name": result.from_name,
                    "to_name": result.to_name,
                    "country_code": result.country_code,
                    "status": result.status,
                    "message": result.message,
                    "network_length_m": result.network_length_m,
                    "metric_distance_m": result.metric_distance_m,
                    "from_snap_distance_m": result.from_snap_distance_m,
                    "to_snap_distance_m": result.to_snap_distance_m,
                }
            )
    return path


def write_routes(settings: Settings, results: list[RouteResult]) -> Path | None:
    successful = [
        result
        for result in results
        if result.status == "success" and result.geometry is not None
    ]
    if not successful:
        return None

    settings.output_dir.mkdir(parents=True, exist_ok=True)
    path = settings.output_dir / settings.routes_filename
    records: list[dict[str, Any]] = []
    for result in successful:
        records.append(
            {
                "mode": "truck",
                "from_id": str(result.from_id),
                "to_id": str(result.to_id),
                "from_name": result.from_name,
                "to_name": result.to_name,
                "country": result.country_code,
                "network_km": result.network_length_m / 1000.0,
                "metric_km": result.metric_distance_m / 1000.0,
                "from_snap_m": result.from_snap_distance_m,
                "to_snap_m": result.to_snap_distance_m,
                "geometry": result.geometry,
            }
        )

    routes = gpd.GeoDataFrame(records, geometry="geometry", crs=settings.working_crs)
    routes = routes.to_crs(settings.output_crs)
    routes.to_file(path, layer="truck_routes", driver="GPKG")
    return path
