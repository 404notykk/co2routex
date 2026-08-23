from __future__ import annotations

from pathlib import Path

import geopandas as gpd
import pandas as pd
from openpyxl import load_workbook
from pyproj import CRS, Transformer
from shapely.geometry import Point
from shapely.strtree import STRtree

from .models import RailwayRequestSettings


REQUIRED_NODE_COLUMNS = {
    "node_id",
    "node_name",
    "longitude",
    "latitude",
    "altitude",
    "annual_flux",
    "node_type",
}


def _read_nodes(workbook_path: Path) -> pd.DataFrame:
    nodes = pd.read_excel(workbook_path, sheet_name="nodes", dtype={"node_id": str})
    missing = REQUIRED_NODE_COLUMNS.difference(nodes.columns)
    if missing:
        raise ValueError(f"The nodes sheet is missing columns: {sorted(missing)}")
    duplicates = nodes.loc[nodes["node_id"].duplicated(), "node_id"].tolist()
    if duplicates:
        raise ValueError(f"Duplicate node IDs in nodes sheet: {duplicates}")
    return nodes


def _read_stations(settings: RailwayRequestSettings) -> gpd.GeoDataFrame:
    kwargs = {"layer": settings.stations_layer} if settings.stations_layer else {}
    stations = gpd.read_file(settings.stations_path, **kwargs)
    if stations.crs is None:
        raise ValueError("The railway station layer must have a defined CRS.")
    for field in (settings.station_id_field, settings.station_name_field):
        if field not in stations.columns:
            raise ValueError(f"Station layer is missing field '{field}'.")
    stations = stations[
        stations.geometry.notna() & ~stations.geometry.is_empty
    ].copy()
    stations["station_id"] = stations[settings.station_id_field].astype(str)
    stations["station_name"] = stations[settings.station_name_field].astype(str)
    if stations["station_id"].duplicated().any():
        duplicates = stations.loc[
            stations["station_id"].duplicated(), "station_id"
        ].tolist()
        raise ValueError(f"Duplicate station IDs: {duplicates}")
    stations.geometry = stations.geometry.map(
        lambda geometry: geometry
        if isinstance(geometry, Point)
        else geometry.representative_point()
    )
    return stations


def _select_nearest_stations(
    nodes: pd.DataFrame,
    stations: gpd.GeoDataFrame,
    nodes_crs: str,
    maximum_distance_m: float,
    station_node_prefix: str,
) -> pd.DataFrame:
    transformer = Transformer.from_crs(
        nodes_crs,
        CRS.from_user_input(stations.crs),
        always_xy=True,
    )
    station_points = list(stations.geometry)
    tree = STRtree(station_points)
    records: list[dict] = []
    source_nodes = nodes[
        ~nodes["node_id"].astype(str).str.startswith(station_node_prefix)
    ]
    for node in source_nodes.itertuples(index=False):
        x, y = transformer.transform(float(node.longitude), float(node.latitude))
        node_point = Point(x, y)
        station_index = int(tree.nearest(node_point))
        station = stations.iloc[station_index]
        proximity_m = float(node_point.distance(station.geometry))
        if proximity_m <= maximum_distance_m:
            records.append(
                {
                    "source_node_id": str(node.node_id),
                    "station_id": str(station.station_id),
                    "station_name": str(station.station_name),
                    "proximity_distance_km": proximity_m / 1000.0,
                }
            )
    return pd.DataFrame(
        records,
        columns=[
            "source_node_id",
            "station_id",
            "station_name",
            "proximity_distance_km",
        ],
    )


def _matrix_labels(worksheet) -> tuple[dict[str, int], dict[str, int]]:
    columns: dict[str, int] = {}
    rows: dict[str, int] = {}
    for column in range(2, worksheet.max_column + 1):
        value = worksheet.cell(row=1, column=column).value
        if value is not None:
            columns[str(value)] = column
    for row in range(2, worksheet.max_row + 1):
        value = worksheet.cell(row=row, column=1).value
        if value is not None:
            rows[str(value)] = row
    return rows, columns


def _expand_matrix(worksheet, node_ids: list[str]) -> tuple[dict[str, int], dict[str, int]]:
    rows, columns = _matrix_labels(worksheet)
    if set(rows) != set(columns):
        raise ValueError(f"Sheet '{worksheet.title}' is not a square node matrix.")
    for node_id in node_ids:
        if node_id not in columns:
            new_column = worksheet.max_column + 1
            worksheet.cell(row=1, column=new_column, value=node_id)
            columns[node_id] = new_column
        if node_id not in rows:
            new_row = worksheet.max_row + 1
            worksheet.cell(row=new_row, column=1, value=node_id)
            rows[node_id] = new_row
    for row_number in rows.values():
        for column_number in columns.values():
            cell = worksheet.cell(row=row_number, column=column_number)
            if cell.value is None:
                cell.value = 0
    return rows, columns


def _reset_matrix(worksheet, rows: dict[str, int], columns: dict[str, int]) -> None:
    for row_number in rows.values():
        for column_number in columns.values():
            worksheet.cell(row=row_number, column=column_number, value=0)


def _append_station_nodes(
    workbook,
    existing_nodes: pd.DataFrame,
    selected_stations: gpd.GeoDataFrame,
    settings: RailwayRequestSettings,
) -> list[str]:
    worksheet = workbook["nodes"]
    headers = {
        str(worksheet.cell(row=1, column=column).value): column
        for column in range(1, worksheet.max_column + 1)
    }
    missing = REQUIRED_NODE_COLUMNS.difference(headers)
    if missing:
        raise ValueError(f"The nodes sheet is missing columns: {sorted(missing)}")
    transformer = Transformer.from_crs(
        selected_stations.crs,
        settings.nodes_crs,
        always_xy=True,
    )
    existing_ids = set(existing_nodes["node_id"].astype(str))
    station_node_ids: list[str] = []
    for station in selected_stations.itertuples(index=False):
        station_node_id = f"rail_{settings.country_code}_{station.station_id}"
        station_node_ids.append(station_node_id)
        if station_node_id in existing_ids:
            continue
        longitude, latitude = transformer.transform(
            station.geometry.x,
            station.geometry.y,
        )
        values = {
            "node_id": station_node_id,
            "node_name": station.station_name,
            "longitude": longitude,
            "latitude": latitude,
            "altitude": settings.station_altitude,
            "annual_flux": 0,
            "node_type": settings.station_node_type,
            "country": settings.country_code,
        }
        new_row = worksheet.max_row + 1
        for field, value in values.items():
            if field in headers:
                worksheet.cell(row=new_row, column=headers[field], value=value)
        existing_ids.add(station_node_id)
    return station_node_ids


def _distance_lookup(distances: pd.DataFrame) -> dict[frozenset[str], float]:
    required = {"from_station_id", "to_station_id", "distance_km"}
    missing = required.difference(distances.columns)
    if missing:
        raise ValueError(f"Station distance table is missing columns: {sorted(missing)}")
    lookup: dict[frozenset[str], float] = {}
    for row in distances.itertuples(index=False):
        lookup[frozenset((str(row.from_station_id), str(row.to_station_id)))] = float(
            row.distance_km
        )
    return lookup


def process_railway_request(settings: RailwayRequestSettings) -> dict:
    if not settings.workbook_path.exists():
        raise FileNotFoundError(f"Input workbook not found: {settings.workbook_path}")
    if not settings.requested:
        return {
            "requested": False,
            "workbook_path": settings.workbook_path,
            "selected_station_count": 0,
        }
    if not settings.stations_path.exists():
        raise FileNotFoundError(f"Railway stations GeoPackage not found: {settings.stations_path}")
    if not settings.distances_path.exists():
        raise FileNotFoundError(f"Station distance table not found: {settings.distances_path}")

    nodes = _read_nodes(settings.workbook_path)
    stations = _read_stations(settings)
    selections = _select_nearest_stations(
        nodes,
        stations,
        settings.nodes_crs,
        settings.maximum_search_radius_km * 1000.0,
        f"rail_{settings.country_code}_",
    )
    selected_ids = list(dict.fromkeys(selections["station_id"].tolist()))
    selected_stations = stations[
        stations["station_id"].isin(selected_ids)
    ].copy()
    if selected_ids:
        selected_stations = (
            selected_stations.set_index("station_id").loc[selected_ids].reset_index()
        )

    workbook = load_workbook(settings.workbook_path)
    required_sheets = {"nodes", *settings.matrix_sheets}
    missing_sheets = required_sheets.difference(workbook.sheetnames)
    if missing_sheets:
        raise ValueError(f"Workbook is missing sheets: {sorted(missing_sheets)}")

    station_node_ids = _append_station_nodes(
        workbook,
        nodes,
        selected_stations,
        settings,
    )
    all_node_ids = nodes["node_id"].astype(str).tolist()
    all_node_ids.extend(
        node_id for node_id in station_node_ids if node_id not in set(all_node_ids)
    )
    matrix_maps: dict[str, tuple[dict[str, int], dict[str, int]]] = {}
    for sheet_name in settings.matrix_sheets:
        matrix_maps[sheet_name] = _expand_matrix(workbook[sheet_name], all_node_ids)

    railway_sheet = workbook["railway"]
    railway_rows, railway_columns = matrix_maps["railway"]
    if settings.reset_railway_matrix:
        _reset_matrix(railway_sheet, railway_rows, railway_columns)

    lookup = _distance_lookup(
        pd.read_csv(
            settings.distances_path,
            dtype={"from_station_id": str, "to_station_id": str},
        )
    )
    inserted_pairs = 0
    for index, from_station_id in enumerate(selected_ids):
        for to_station_id in selected_ids[index + 1 :]:
            distance_km = lookup.get(frozenset((from_station_id, to_station_id)))
            if distance_km is None:
                continue
            from_node = f"rail_{settings.country_code}_{from_station_id}"
            to_node = f"rail_{settings.country_code}_{to_station_id}"
            railway_sheet.cell(
                row=railway_rows[from_node],
                column=railway_columns[to_node],
                value=distance_km,
            )
            railway_sheet.cell(
                row=railway_rows[to_node],
                column=railway_columns[from_node],
                value=distance_km,
            )
            inserted_pairs += 1

    workbook.save(settings.workbook_path)
    return {
        "requested": True,
        "workbook_path": settings.workbook_path,
        "selected_station_count": len(selected_ids),
        "inserted_station_pairs": inserted_pairs,
        "maximum_search_radius_km": settings.maximum_search_radius_km,
        "selections": selections.to_dict(orient="records"),
    }

