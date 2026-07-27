"""Create an Excel input workbook with the required routing layout."""

from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


NODE_COLUMNS = [
    "node_id",
    "node_name",
    "longitude",
    "latitude",
    "altitude",
    "annual_flux",
    "node_type",
    "country_code",
]
MODES = ("pipeline", "truck", "railway")

HEADER_FILL = PatternFill("solid", fgColor="D9EAD3")
DIAGONAL_FILL = PatternFill("solid", fgColor="E7E6E6")
HEADER_FONT = Font(bold=True)
DEFAULT_COLUMN_WIDTH = 8.83203125
BASE_COLUMN_WIDTH = 10


def configure_default_dimensions(sheet) -> None:
    """Apply the workbook's default column sizing."""
    sheet.sheet_format.defaultColWidth = DEFAULT_COLUMN_WIDTH
    sheet.sheet_format.baseColWidth = BASE_COLUMN_WIDTH


def create_template(path: Path, node_count: int) -> None:
    """Create a routing-input workbook matching ``node_metrics_template.xlsx``."""
    if node_count < 1:
        raise ValueError("node_count must be at least 1")

    path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()

    nodes = workbook.active
    nodes.title = "nodes"
    configure_default_dimensions(nodes)
    nodes.append(NODE_COLUMNS)

    for index in range(1, node_count + 1):
        nodes.append(
            [
                index,
                f"Node {index}",
                None,
                None,
                None,
                None,
                None,
                None,
            ]
        )

    for cell in nodes[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL

    nodes.freeze_panes = "A2"
    nodes.sheet_view.zoomScale = 171

    # These widths reproduce the supplied workbook.
    column_widths = {
        "A": 14,
        "B": 24,
        "C": 16,
        "D": 16,
        "E": 14,
        "F": 16,
        "G": 16,
        "H": 15.5,
    }
    for column, width in column_widths.items():
        nodes.column_dimensions[column].width = width

    for mode in MODES:
        sheet = workbook.create_sheet(mode)
        configure_default_dimensions(sheet)

        # A1 is intentionally blank. Node IDs form the top and left headers.
        for index in range(1, node_count + 1):
            sheet.cell(row=1, column=index + 1, value=index)
            sheet.cell(row=index + 1, column=1, value=index)

        for source in range(1, node_count + 1):
            for target in range(1, node_count + 1):
                sheet.cell(row=source + 1, column=target + 1, value=0)

        for cell in sheet[1]:
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL

        # Style A1 as part of the header band, even though it is blank.
        sheet["A1"].font = HEADER_FONT
        sheet["A1"].fill = HEADER_FILL

        for row in range(2, node_count + 2):
            sheet.cell(row=row, column=1).font = HEADER_FONT
            sheet.cell(row=row, column=row).fill = DIAGONAL_FILL

        sheet.freeze_panes = "B2"

    workbook.save(path)


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Create the node and transport-mode routing workbook."
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("node_metrics_template.xlsx"),
        help="Output workbook path (default: node_metrics_template.xlsx).",
    )
    parser.add_argument(
        "--nodes",
        type=int,
        default=40,
        help="Number of nodes represented in the workbook (default: 40).",
    )
    args = parser.parse_args()

    create_template(args.output, args.nodes)
    print(f"Created {args.output}")


if __name__ == "__main__":
    main()