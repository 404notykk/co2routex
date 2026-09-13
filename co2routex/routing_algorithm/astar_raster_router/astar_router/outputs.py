"""Write a routed workbook copy, GIS routes, and one benchmark workbook."""
import json
import os
from pathlib import Path
from uuid import uuid4
import fiona
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from .inputs import _clean_id

RESULT_FIELDS = ["run_id", "pair_id", "mode", "from_id", "to_id", "from_name", "to_name",
                 "status", "message", "result_source", "straight_line_km", "distance_km",
                 "accumulated_resistance", "average_route_resistance", "path_cells",
                 "explored_cells", "discovered_cells", "queue_peak_entries",
                 "from_cell_center_offset_m", "to_cell_center_offset_m"]
TIME_FIELDS = ["run_id", "pair_id", "from_id", "to_id", "pair_setup_wall_s",
               "search_setup_wall_s", "search_wall_s", "path_build_wall_s",
               "route_build_wall_s", "pair_other_wall_s", "pair_total_wall_s"]
MEMORY_FIELDS = ["run_id", "pair_id", "from_id", "to_id", "pair_start_rss_mib", "pair_peak_rss_mib"]


def scalar(value):
    if isinstance(value, Path): return str(value)
    if isinstance(value, (dict, list, tuple)): return json.dumps(value, default=str)
    return value


def table(wb, title, fields, records):
    ws = wb.create_sheet(title)
    ws.append(fields)
    for record in records:
        ws.append([scalar(record.get(field)) for field in fields])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_view.showGridLines = False
    for cell in ws[1]:
        cell.font = Font(name="Calibri", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="24465B")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 42
    for col, field in enumerate(fields, 1):
        width = 24
        if field in ("description", "value", "message"): width = 65
        if field in ("field", "metric"): width = 38
        if field.endswith("_id"): width = 16
        ws.column_dimensions[ws.cell(1,col).column_letter].width = width
        for row in range(2, ws.max_row+1):
            cell = ws.cell(row,col)
            cell.font = Font(name="Calibri", size=11)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if isinstance(cell.value, str): cell.data_type = "s"
            if isinstance(cell.value, float): cell.number_format = "0.000000" if field.endswith("_s") else "0.000"
            if row % 2 == 0: cell.fill = PatternFill("solid", fgColor="EFF4F7")
    return ws


def save(wb, path):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.stem + f".pending-{uuid4().hex}.xlsx")
    try:
        wb.save(temporary)
        os.replace(temporary, path)
    finally:
        wb.close()
        temporary.unlink(missing_ok=True)


def write_routed_workbook(settings, records):
    target = settings.output_dir / settings.output_workbook_name
    if target.resolve() == settings.workbook_path.resolve():
        raise ValueError("Output must not overwrite the input workbook")
    wb = load_workbook(settings.workbook_path)
    try:
        for mode in settings.mode_sheets:
            ws = wb[mode]
            cols = {_clean_id(ws.cell(1,c).value): c for c in range(2,ws.max_column+1)
                    if ws.cell(1,c).value is not None}
            rows = {_clean_id(ws.cell(r,1).value): r for r in range(2,ws.max_row+1)
                    if ws.cell(r,1).value is not None}
            for record in records:
                if record["mode"] == mode and record["status"] == "ok":
                    cell = ws.cell(rows[record["from_id"]],cols[record["to_id"]])
                    cell.value = record["distance_km"]
                    cell.number_format = "0.000"
        if "pipeline_route_metrics" in wb: del wb["pipeline_route_metrics"]
        table(wb, "pipeline_route_metrics", RESULT_FIELDS, records)
        save(wb, target)
    finally:
        wb.close()


def write_geopackage(settings, records, crs):
    settings.output_dir.mkdir(parents=True, exist_ok=True)
    target = settings.output_dir / settings.routes_filename
    temporary = target.with_name(target.stem + f".pending-{uuid4().hex}.gpkg")
    properties = {"mode":"str", "from_id":"str", "to_id":"str", "from_name":"str", "to_name":"str",
                  "status":"str", "distance_km":"float", "accumulated_resistance":"float",
                  "average_route_resistance":"float", "explored_cells":"int", "path_cells":"int",
                  "from_cell_center_offset_m":"float", "to_cell_center_offset_m":"float"}
    try:
        with fiona.open(temporary,"w",driver="GPKG",layer="pipeline_routes",
                        schema={"geometry":"LineString","properties":properties},crs_wkt=crs.to_wkt()) as dst:
            for record in records:
                if record["status"] != "ok": continue
                coords = record["coordinates"]
                if len(coords) == 1: coords = [coords[0],coords[0]]
                dst.write({"geometry":{"type":"LineString","coordinates":coords},
                           "properties":{key:record.get(key) for key in properties}})
        os.replace(temporary, target)
    finally:
        temporary.unlink(missing_ok=True)


DEFINITIONS = {
    "total_pairs": "Count of requested directed connections in this run.",
    "successful_pairs": "Count with status ok.",
    "failed_pairs": "Count with status no_path; failed matrix entries retain their original request values.",
    "computed_pairs": "Pairs that invoked A*, including failed searches.",
    "cached_pairs": "Pairs served from direct or reverse route cache.",
    "total_search_wall_s": "Sum of search_wall_s across pairs; already included in all_pairs_wall_s and run_total_wall_s.",
    "workbook_input_wall_s": "Run component: read/validate node workbook and permitted connections, once per run.",
    "raster_read_wall_s": "Run component: open raster, read full Float32 band and validity mask, including decoding; close file.",
    "raster_prepare_wall_s": "Run component: create Boolean mask, validate invariant values and calculate minimum/maximum once.",
    "node_prepare_wall_s": "Run component: transform all node coordinates and reject outside/invalid cells; no snapping.",
    "all_pairs_wall_s": "Inclusive run total for the pair loop; contains pair times, logging and monitor overhead. Do not add pair times to this.",
    "route_workbook_write_wall_s": "Run component: write updated input-workbook copy and route metrics.",
    "gis_write_wall_s": "Run component: write successful route geometries to a GeoPackage.",
    "run_other_wall_s": "Run residual: run_total minus non-overlapping setup, pair-loop and output components.",
    "run_total_wall_s": "Inclusive elapsed run time from input loading through normal outputs and monitor stop. Excludes benchmark report construction/export, config/imports and benchmark aggregation.",
    "pair_setup_wall_s": "Pair component: prepare result identity, compute straight-line distance, and look up optional cached route.",
    "search_setup_wall_s": "Pair component: endpoint checks and search-state initialization; no full-raster validation.",
    "search_wall_s": "Pair component: A* priority-queue exploration only; zero for cache reuse.",
    "path_build_wall_s": "Pair component: parent-chain reconstruction and raster route-length calculation; zero for cache reuse.",
    "route_build_wall_s": "Pair component: calculate route properties and map path cells to coordinates.",
    "pair_other_wall_s": "Residual: pair total minus five timed pair components; includes call/cleanup overhead.",
    "pair_total_wall_s": "Inclusive pair wall time; do not add it to its components. RSS monitor calls are outside this timer.",
    "baseline_rss_mib": "Process RSS after shared raster/node setup, before first pair. MiB = bytes / 1048576.",
    "pair_start_rss_mib": "Process RSS immediately before this pair; includes shared setup and retained earlier results.",
    "pair_peak_rss_mib": "Highest sampled process RSS during this pair; native memory included, short peaks may be missed. Not additive across pairs/stages.",
    "run_peak_rss_mib": "Highest sampled process RSS during this run, including normal output generation; excludes report export.",
    "rss_sample_interval_s": "Requested memory-sampling interval. Thread scheduling/GIL may make actual intervals longer.",
    "explored_cells": "Cells finalized by this search, including the goal; zero for a cache hit (no new exploration).",
    "discovered_cells": "Unique cells with a known cost during this search; zero for cache reuse.",
    "queue_peak_entries": "Maximum heap length; may include stale entries, so not unique pending-cell count.",
    "path_cells": "Cells in the final path; not the number explored.",
    "straight_line_km": "Ellipsoidal geodesic distance between original input coordinates.",
    "distance_km": "Length along the selected raster-cell-center route in projected metres / 1000; original-coordinate offsets excluded.",
    "accumulated_resistance": "Sum over path edges: edge length (metres) times mean endpoint resistance. Not monetary cost.",
    "average_route_resistance": "Accumulated resistance / route length in metres; blank for zero-length routes.",
    "from_cell_center_offset_m": "Distance from input coordinate to its containing cell center; no adjacent-cell snapping.",
    "to_cell_center_offset_m": "Distance from destination coordinate to its containing cell center; no adjacent-cell snapping.",
    "status": "ok or no_path. Invalid node input aborts the run before routing with a node-specific error.",
    "message": "Failure explanation; blank for success.",
    "result_source": "computed, cache or reverse_cache; benchmark configurations disable all route caching.",
    "cache_max_routes": "LRU route-cache capacity; 0 disables storage and reuse. Output coordinate records still accumulate until writing.",
    "search_domain": "full_valid_raster: no per-pair geographical window or buffer.",
    "pairs_per_hour": "3600 * pair count / all_pairs_wall_s; observed throughput for this case set, not a general prediction.",
    "run_id": "Measured repetition identifier; each repetition loads the full raster once.",
    "pair_id": "Connection ordinal within a run; join with run_id across sheets.",
    "width_cells": "Number of raster columns loaded; whole input raster, not a per-pair window.",
    "height_cells": "Number of raster rows loaded; whole input raster, not a per-pair window.",
    "valid_cells": "Finite, unmasked traversable cells; excludes zero if zero_is_barrier is enabled.",
    "minimum_resistance": "Global minimum over traversable cells, calculated once and used in the admissible A* heuristic.",
    "maximum_resistance": "Global maximum over traversable cells, calculated once during validation.",
    "resistance_array_mib": "Exact bytes in the Float32 resistance array / 1048576; not total process RAM.",
    "mask_array_mib": "Exact bytes in the Boolean mask / 1048576; not total process RAM.",
    "compression": "On-disk pixel compression; decoded arrays have the same size for LZW and NONE.",
    "disk_dtype": "Stored pixel data type; this version requires Float32 without silent conversion.",
    "memory_dtype": "Resistance array type retained in RAM: Float32. Cost arithmetic uses Python double-precision floats.",
    "block_rows": "On-disk tile/block height, not a search limit.",
    "block_columns": "On-disk tile/block width, not a search limit.",
    "cell_x_m": "Raster cell resolution along the column direction in metres.",
    "cell_y_m": "Raster cell resolution along the row direction in metres.",
    "file_size_mib": "File size on disk / 1048576; excludes companion files if any.",
    "snap_radius_cells": "Must be 0: outside/invalid containing cells raise an error; no relocation.",
}


def write_benchmark(path, runs, records, metadata):
    wb = Workbook(); del wb[wb.sheetnames[0]]
    table(wb,"Run_summary",list(runs[0]) if runs else ["run_id"],runs)
    table(wb,"Pair_results",RESULT_FIELDS,records)
    table(wb,"Pair_runtime",TIME_FIELDS,records)
    table(wb,"Memory",MEMORY_FIELDS,records)
    metadata_rows = [{"run_id":run_id,"field":key,"value":value}
                     for run_id, values in metadata.items() for key,value in values.items()]
    table(wb,"Settings",["run_id","field","value"],metadata_rows)
    table(wb,"Field_guide",["field","description"],
          [{"field":key,"description":value} for key,value in DEFINITIONS.items()])
    save(wb,path)
