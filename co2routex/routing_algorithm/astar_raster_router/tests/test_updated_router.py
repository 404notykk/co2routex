"""Correctness checks: independent Dijkstra, barriers, I/O and timing scopes."""
import heapq
import math
from pathlib import Path
from unittest.mock import patch
import numpy as np
import pandas as pd
import pytest
import rasterio
from rasterio.transform import from_origin
from pyproj import Transformer
from openpyxl import load_workbook
from astar_router.astar import PreparedGrid, astar_search, search_prepared, NoPathError
from astar_router.raster import ResistanceRaster
from astar_router.models import Settings, Node
from astar_router.pipeline import run_with_timing


def dijkstra(cost, valid, start, goal):
    """Independent small-grid oracle, fixed 100 m square cells and no corner cutting."""
    distance = {start:0.0}; queue = [(0.0,start)]
    while queue:
        accumulated, cell = heapq.heappop(queue)
        if accumulated != distance[cell]: continue
        if cell == goal: return accumulated
        r,c = cell
        for dr in (-1,0,1):
            for dc in (-1,0,1):
                nr,nc = r+dr,c+dc
                if (dr,dc)==(0,0) or not (0<=nr<cost.shape[0] and 0<=nc<cost.shape[1]): continue
                if not valid[nr,nc]: continue
                if dr and dc and not (valid[r+dr,c] and valid[r,c+dc]): continue
                candidate = accumulated+100*math.hypot(dr,dc)*(float(cost[r,c])+float(cost[nr,nc]))/2
                if candidate < distance.get((nr,nc),math.inf):
                    distance[nr,nc]=candidate;heapq.heappush(queue,(candidate,(nr,nc)))
    return math.inf


@pytest.mark.parametrize("seed",range(8))
def test_exact_cost_against_independent_oracle(seed):
    rng=np.random.default_rng(seed)
    values=rng.uniform(0,60,(10,11)).astype("float32")
    valid=rng.random(values.shape)>0.15;valid[0,0]=valid[-1,-1]=True
    expected=dijkstra(values,valid,(0,0),(9,10))
    if math.isinf(expected):
        with pytest.raises(NoPathError): astar_search(values,valid,from_origin(0,1000,100,100),(0,0),(9,10))
    else:
        result=astar_search(values,valid,from_origin(0,1000,100,100),(0,0),(9,10))
        assert math.isclose(result.accumulated_resistance,expected,rel_tol=1e-12)
        assert all(valid[cell] for cell in result.path)


def test_barriers_and_queue_insertions():
    values=np.ones((5,5),dtype="float32");values[2,2]=-9999
    valid=values!=-9999
    original=heapq.heappush
    def checked(queue,item):
        assert valid[item[2]]
        return original(queue,item)
    with patch("astar_router.astar.heapq.heappush",checked):
        astar_search(values,valid,from_origin(0,500,100,100),(2,0),(2,4))
    with pytest.raises(ValueError,match="not traversable"):
        astar_search(values,valid,from_origin(0,500,100,100),(2,2),(2,4))
    mask=np.array([[True,False],[False,True]])
    with pytest.raises(NoPathError):
        astar_search(np.ones((2,2)),mask,from_origin(0,200,100,100),(0,0),(1,1))


def write_raster(path,compression="lzw",barrier=False):
    data=np.full((32,32),2.0025,dtype="float32")
    data[0,:]=-9999
    if barrier: data[16,:]=-9999
    transform=from_origin(4321000,3210000,100,100)
    with rasterio.open(path,"w",driver="GTiff",width=32,height=32,count=1,dtype="float32",
                       crs="EPSG:3035",transform=transform,nodata=-9999,tiled=True,
                       blockxsize=128,blockysize=128,compress=compression) as dst: dst.write(data,1)
    return transform


def write_nodes(path,transform):
    projection=Transformer.from_crs(3035,4326,always_xy=True)
    records=[]
    for nid,row,col in (("1",3,3),("2",27,27),("3",5,27)):
        x,y=transform*(col+0.5,row+0.5);lon,lat=projection.transform(x,y)
        records.append(dict(node_id=nid,node_name=f"Node {nid}",longitude=lon,latitude=lat,
                            altitude=10,annual_flux=1,node_type="emitter",country_code="TEST"))
    matrix=pd.DataFrame([[0,1,1],[1,0,0],[0,0,0]],index=["1","2","3"],columns=["1","2","3"])
    with pd.ExcelWriter(path) as writer:
        pd.DataFrame(records).to_excel(writer,sheet_name="nodes",index=False)
        matrix.to_excel(writer,sheet_name="pipeline")


def test_float32_mask_and_strict_nodes(tmp_path):
    p=tmp_path/"r.tif";tr=write_raster(p);r=ResistanceRaster(p)
    assert r.resistance.dtype==np.float32 and r.grid.minimum==float(np.float32(2.0025))
    assert not r.traversable[0,0] and r.resistance.flags.writeable is False
    transformer=Transformer.from_crs(3035,4326,always_xy=True)
    lon,lat=transformer.transform(*(tr*(0.5,0.5)))
    with pytest.raises(ValueError,match="no snapping"):
        r.node_cells({"bad":Node("bad",lon,lat)},"EPSG:4326",0)
    lon,lat=transformer.transform(*(tr*(-10,0.5)))
    with pytest.raises(ValueError,match="outside"):
        r.node_cells({"out":Node("out",lon,lat)},"EPSG:4326",0)
    with pytest.raises(ValueError,match="non-negative"):
        PreparedGrid(np.array([[-1.]],dtype="float32"),np.ones((1,1),bool),tr)


def test_full_run_and_timing_reconciliation(tmp_path):
    p=tmp_path/"r.tif";x=tmp_path/"nodes.xlsx";tr=write_raster(p);write_nodes(x,tr)
    settings=Settings(p,x,tmp_path/"out")
    original=PreparedGrid.__init__;calls=[]
    def counted(self,*args,**kwargs): calls.append(1);original(self,*args,**kwargs)
    with patch.object(PreparedGrid,"__init__",counted):
        records,summary,metadata=run_with_timing(settings)
    assert len(calls)==1 and len(records)==3 and summary["failed_pairs"]==0
    assert summary["cached_pairs"]==0
    components=("pair_setup_wall_s","search_setup_wall_s","search_wall_s","path_build_wall_s",
                "route_build_wall_s","pair_other_wall_s")
    for record in records:
        assert all(record[k]>=0 for k in components)
        assert abs(sum(record[k] for k in components)-record["pair_total_wall_s"])<1e-9
        assert record["pair_peak_rss_mib"]>=record["pair_start_rss_mib"]>0
    run_components=("workbook_input_wall_s","raster_read_wall_s","raster_prepare_wall_s",
                    "node_prepare_wall_s","all_pairs_wall_s","route_workbook_write_wall_s",
                    "gis_write_wall_s","run_other_wall_s")
    assert abs(sum(summary[k] for k in run_components)-summary["run_total_wall_s"])<1e-9
    assert metadata["memory_dtype"]=="float32"
    wb=load_workbook(settings.output_dir/"benchmark_results.xlsx",data_only=True)
    assert wb.sheetnames==["Run_summary","Pair_results","Pair_runtime","Memory","Settings","Field_guide"]
    assert wb["Pair_runtime"].max_row==4
    wb.close()
    assert not list(settings.output_dir.glob("*.csv"))
    wb=load_workbook(settings.output_dir/"node_metrics_routed.xlsx",data_only=True)
    assert wb["pipeline"].cell(2,3).value==records[0]["distance_km"]
    wb.close()


def test_compression_and_cache_do_not_change_routes(tmp_path):
    a=tmp_path/"a.tif";b=tmp_path/"b.tif";x=tmp_path/"nodes.xlsx"
    tr=write_raster(a);write_raster(b,"NONE");write_nodes(x,tr)
    r1,_,_=run_with_timing(Settings(a,x,tmp_path/"one"),write_timing_files=False)
    r2,s2,_=run_with_timing(Settings(b,x,tmp_path/"two",cache_max_routes=4,cache_reverse_routes=True),write_timing_files=False)
    assert s2["cached_pairs"]==1
    assert [r["distance_km"] for r in r1]==[r["distance_km"] for r in r2]
    assert [r["accumulated_resistance"] for r in r1]==[r["accumulated_resistance"] for r in r2]
    assert next(r for r in r2 if r["result_source"]=="reverse_cache")["search_wall_s"]==0


def test_no_path_report(tmp_path):
    p=tmp_path/"r.tif";x=tmp_path/"nodes.xlsx";tr=write_raster(p,barrier=True);write_nodes(x,tr)
    rows,summary,_=run_with_timing(Settings(p,x,tmp_path/"out"))
    assert summary["failed_pairs"]==2
    assert all(r["message"] for r in rows if r["status"]=="no_path")


def test_variant_writer_preserves_pixels_and_mask(tmp_path):
    from prepare_raster_variants import prepare
    source=tmp_path/"source.tif"; target=tmp_path/"uncompressed.tif"
    write_raster(source)
    prepare(source,target)
    prepare(source,target)  # Existing verified output is reusable.
    with rasterio.open(source) as src, rasterio.open(target) as dst:
        assert dst.compression is None and dst.block_shapes==[(128,128)]
        np.testing.assert_array_equal(src.read(1),dst.read(1))
        np.testing.assert_array_equal(src.read_masks(1),dst.read_masks(1))


def test_benchmark_entrypoint_with_relative_yaml(tmp_path, monkeypatch):
    import sys
    import yaml
    from run_distance_benchmark import main
    p=tmp_path/"r.tif"; x=tmp_path/"nodes.xlsx"
    tr=write_raster(p); write_nodes(x,tr)
    config=tmp_path/"benchmark.yaml"
    config.write_text(yaml.safe_dump(dict(raster_path="r.tif",workbook_path="nodes.xlsx",
                         output_dir="report",benchmark=dict(warmup_runs=0,repetitions=1))))
    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(sys,"argv",["run_distance_benchmark.py","--config",str(config)])
    assert main()==0
    assert (tmp_path/"report"/"benchmark_results.xlsx").is_file()
