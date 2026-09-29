"""Parquet/GeoParquet mapping and physical row-selection regression checks."""
from __future__ import annotations
import csv
import hashlib
import json
import tempfile
import unittest
from dataclasses import replace
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook
from pipeline_cost_model import PipelineCostConfig
from pipeline_cost_model.inputs import read_routes
from pipeline_cost_model.runner import run


class ParquetInputTests(unittest.TestCase):
    def make_source(self, directory):
        source = Path(directory) / "routes.parquet"
        frame = pd.DataFrame({
            "emitter_name": ["Not pipeline", "Selected emitter", "Invalid unselected flow"],
            "sink_name": ["Aramis"] * 3,
            "emitter_id": ["000", "001", "003"], "sink_id": ["002"] * 3,
            "distance_km": [10.0, 200.0, 20.0],
            "co2_t_per_yr": [-1.0, 1_000_000.0, -5.0],
            "average_route_resistance": [1.0, 1.2, 1.0],
            "mode": ["truck", "pipeline", "pipeline"], "country": ["NL"] * 3,
            "bbox.xmin": [100, 200, 300], "bbox": [{"xmin": 100}] * 3,
            # Ignored binary values must not enter the engineering or Excel code.
            "geometry": [b"\xff\x00"] * 3,
        }, index=pd.Index([900, 12, 12], name="saved_index"))
        frame.to_parquet(source, engine="pyarrow", index=True, row_group_size=1)
        config = PipelineCostConfig(
            input_path=source, output_dir=Path(directory)/"out", mode="benchmark",
            from_id_column="emitter_name", to_id_column="sink_name",
            annual_flow_column="co2_t_per_yr", input_row=2,
        )
        return source, config

    def test_parquet_selected_row_with_saved_index_and_name_or_id_mapping(self):
        with tempfile.TemporaryDirectory() as d:
            source, config = self.make_source(d)
            config.validate()
            routes = read_routes(config)
            self.assertEqual(len(routes), 1)
            route = routes[0]
            self.assertEqual((route.from_id, route.to_id), ("Selected emitter", "Aramis"))
            self.assertEqual(route.annual_flow_t_per_year, 1_000_000)
            self.assertEqual(route.average_route_resistance, 1.2)
            self.assertEqual(route.source_row, 3)  # Excel-style row, including header.
            by_id = read_routes(replace(config, from_id_column="emitter_id", to_id_column="sink_id"))[0]
            self.assertEqual((by_id.from_id, by_id.to_id), ("001", "002"))
            # Selection must happen before validation of other routes.
            with self.assertRaisesRegex(ValueError, "positive"):
                read_routes(replace(config, input_row=None))
            renamed = source.with_suffix(".geoparquet")
            source.rename(renamed)
            alternative = replace(config, input_path=renamed)
            alternative.validate()
            self.assertEqual(read_routes(alternative)[0], route)

    def test_parquet_both_mode_exports_selected_costs_and_safe_attributes(self):
        with tempfile.TemporaryDirectory() as d:
            source, config = self.make_source(d)
            original_hash = hashlib.sha256(source.read_bytes()).hexdigest()
            config = replace(config, mode="both", annual_flow_min_t_per_year=500_000,
                             annual_flow_max_t_per_year=2_000_000, validation_samples=2)
            outputs = run(config)
            self.assertEqual(set(outputs), {config.output_workbook_name, "routes_costed.parquet"})
            original = pd.read_parquet(source)
            copied = pd.read_parquet(outputs["routes_costed.parquet"])
            pd.testing.assert_frame_equal(copied[original.columns], original)
            self.assertTrue(pd.isna(copied.iloc[0]["pipeline_cost_basis"]))
            self.assertEqual(copied.iloc[1]["pipeline_cost_basis"], "baseline")
            self.assertTrue(pd.isna(copied.iloc[2]["pipeline_annual_cost_eur_per_year"]))
            self.assertAlmostEqual(copied.iloc[1]["pipeline_annual_cost_eur_per_year"], 15463744.110506363, places=5)
            workbook = load_workbook(outputs[config.output_workbook_name])
            sheet = workbook["summary"]
            headers = {cell.value: cell.column for cell in sheet[1]}
            self.assertEqual(sheet.max_row, 2)
            self.assertNotIn("geometry", headers)
            self.assertEqual(sheet.cell(2, headers["from_id"]).value, "Selected emitter")
            self.assertEqual(sheet.cell(2, headers["input_data_row"]).value, 2)
            workbook.close()
            self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest(), original_hash)

    def test_row_errors_do_not_substitute_another_route(self):
        with tempfile.TemporaryDirectory() as d:
            _, config = self.make_source(d)
            with self.assertRaisesRegex(ValueError, "excluded.*no other row"):
                read_routes(replace(config, input_row=1))
            with self.assertRaisesRegex(ValueError, "out of range.*3 data rows"):
                read_routes(replace(config, input_row=4))
            for invalid in (0, -1, True, 1.5):
                with self.subTest(invalid=invalid), self.assertRaisesRegex(ValueError, "input_row"):
                    replace(config, input_row=invalid).validate()
            with self.assertRaisesRegex(ValueError, "Missing input columns.*not_present"):
                read_routes(replace(config, from_id_column="not_present"))
            with self.assertRaisesRegex(ValueError, "benchmark requires.*co2_per_yr"):
                read_routes(replace(config, annual_flow_column="co2_per_yr"))

    def test_csv_and_excel_use_the_same_mapping_and_physical_row_selection(self):
        with tempfile.TemporaryDirectory() as d:
            frame = pd.DataFrame({"origin": ["skip", "001"], "destination": ["000", "002"],
                                  "length": [20, 200], "mass": [-1, 1_000_000], "weight": [1, 1.2],
                                  "mode": ["truck", "pipeline"]})
            for extension in ("csv", "xlsx"):
                with self.subTest(extension=extension):
                    source = Path(d)/f"input.{extension}"
                    if extension == "csv":
                        frame.to_csv(source, index=False)
                    else:
                        frame.to_excel(source, index=False, sheet_name="pipeline_route_metrics")
                    config = PipelineCostConfig(input_path=source, output_dir=Path(d)/extension, mode="benchmark",
                        from_id_column="origin", to_id_column="destination", distance_column="length",
                        annual_flow_column="mass", resistance_column="weight", input_row=2)
                    route = read_routes(config)[0]
                    self.assertEqual((route.from_id, route.to_id), ("001", "002"))
                    self.assertEqual(route.source_row, 3)
                    outputs = run(config)
                    key = "pipeline_annual_cost_eur_per_year"
                    if extension == "xlsx":
                        workbook = load_workbook(outputs["input_costed.xlsx"])
                        sheet = workbook[config.route_metrics_sheet]
                        headers = {cell.value: cell.column for cell in sheet[1]}
                        self.assertEqual(sheet.max_row, 3)
                        self.assertIsNone(sheet.cell(2, headers[key]).value)
                        self.assertIsNotNone(sheet.cell(3, headers[key]).value)
                        self.assertEqual(sheet.cell(3, headers["origin"]).value, "001")
                        workbook.close()
                    else:
                        with outputs["input_costed.csv"].open(encoding="utf-8-sig") as stream:
                            rows = list(csv.DictReader(stream))
                        self.assertEqual(len(rows), 2)
                        self.assertEqual(rows[0][key], "")
                        self.assertGreater(float(rows[1][key]), 0)
                        self.assertEqual(rows[1]["origin"], "001")



if __name__ == "__main__":
    unittest.main()
