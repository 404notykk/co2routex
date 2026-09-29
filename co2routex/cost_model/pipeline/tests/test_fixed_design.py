"""Accounting coefficients at one reference flow, without range fitting."""
from dataclasses import replace
import json
from pathlib import Path
import tempfile
import unittest
import pandas as pd
from openpyxl import load_workbook
from pipeline_cost_model import (
    PipelineCostConfig, PipelineCostModel, RouteInput, FixedDesignApproximation,
    calculate_pipeline_cost,
)
from pipeline_cost_model.annual_costs import COMPONENT_FIELDS
from pipeline_cost_model.runner import run


class FixedDesignTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.config = PipelineCostConfig(mode="both", coefficient_method="fixed_design",
            annual_flow_min_t_per_year=0, annual_flow_max_t_per_year=0,
            fit_method="not_used", fit_samples=0, validation_samples=0, fit_error_warning_percent=0)
        cls.route = RouteInput("A", "B", 200, 1.7, "reference", 1_000_000)
        cls.model = PipelineCostModel(cls.config)
        cls.result = cls.model.calculate_route(cls.route)

    def test_reference_total_and_components_reconcile_without_any_fit_settings(self):
        a = self.result.approximation
        self.assertIsInstance(a, FixedDesignApproximation)
        self.assertIs(a.reference, self.result.benchmark)
        for basis, components in a.coefficients.items():
            costs = self.result.benchmark.costs[basis]
            for name, pair in components.items():
                self.assertAlmostEqual(pair.value(1_000_000), costs[COMPONENT_FIELDS[name]], places=6)
            self.assertEqual(components["annualized_capex"].slope_eur_per_t, 0)
            self.assertEqual(components["fixed_opex"].slope_eur_per_t, 0)
            self.assertEqual(components["electricity_opex"].intercept_eur_per_year, 0)
            self.assertAlmostEqual(components["total"].intercept_eur_per_year,
                costs["annualized_capex_total_eur_per_year"] + costs["fixed_opex_total_eur_per_year"], places=6)
        pair = a.coefficients["baseline"]["total"]
        self.assertAlmostEqual(pair.intercept_eur_per_year, 11906100.565173205, places=5)
        self.assertAlmostEqual(pair.slope_eur_per_t, 3.5576435453331564, places=10)
        self.assertAlmostEqual(pair.intercept_eur_per_year/1_000_000 + pair.slope_eur_per_t,
                               a.reference.costs["baseline"]["levelized_cost_eur_per_t"], places=10)

    def test_cost_accounting_holds_investment_fixed_and_handles_build_selection(self):
        a = self.result.approximation
        total = a.coefficients["baseline"]["total"]
        self.assertEqual(a.predict(0), total.intercept_eur_per_year)
        self.assertEqual(a.predict(0, built=False), 0)
        self.assertAlmostEqual(a.predict(500_000), total.intercept_eur_per_year + total.slope_eur_per_t*500_000)
        self.assertEqual(a.predict(500_000, component="annualized_capex"),
                         a.predict(1_000_000, component="annualized_capex"))
        for flow in (-1, float("nan"), float("inf"), 1_000_001):
            with self.subTest(flow=flow), self.assertRaises(ValueError):
                a.predict(flow)
        with self.assertRaises(ValueError):
            a.predict(1, built=False)

    def test_two_number_convenience_function_returns_coefficients_and_details(self):
        result = calculate_pipeline_cost(200, 1_000_000)
        self.assertEqual(result.approximation.coefficient_method, "fixed_design")
        self.assertAlmostEqual(result.approximation.predict(1_000_000),
                               result.benchmark.costs["baseline"]["total_annual_cost_eur_per_year"], places=6)
        self.assertEqual(result.benchmark.operating_flow_t_per_h, 125)
        self.assertIn("capex_material_eur", result.benchmark.costs["baseline"])

    def test_reference_flow_is_required_and_global_fallback_works(self):
        route = replace(self.route, annual_flow_t_per_year=None)
        model = PipelineCostModel(replace(self.config, mode="approximation"))
        with self.assertRaisesRegex(ValueError, "require annual transported flow"):
            model.calculate_route(route)
        model = PipelineCostModel(replace(self.config, mode="approximation", benchmark_annual_flow_t_per_year=1_000_000))
        result = model.calculate_route(route)
        self.assertEqual(result.approximation.reference_annual_flow_t_per_year, 1_000_000)
        self.assertIsNone(result.benchmark)

    def test_zero_electricity_remains_a_fixed_design_cost(self):
        settings = replace(self.config, electricity_price_eur_per_mwh=0)
        result = PipelineCostModel(settings).calculate_route(self.route)
        for basis, components in result.approximation.coefficients.items():
            self.assertEqual(components["total"].slope_eur_per_t, 0)
            self.assertAlmostEqual(components["total"].intercept_eur_per_year,
                                   result.benchmark.costs[basis]["total_annual_cost_eur_per_year"], places=6)
        coefficients = result.approximation.coefficients
        self.assertEqual(set(coefficients), {"baseline"})
        self.assertGreater(result.benchmark.energy["total_mwh_per_year"], 0)

    def test_parquet_row_selection_exports_both_structures_without_bounds(self):
        with tempfile.TemporaryDirectory() as d:
            source = Path(d)/"routes.parquet"
            pd.DataFrame({"emitter_name": ["unselected", "A"], "sink_name": ["B", "B"],
                "distance_km": [0, 200], "co2_t_per_yr": [-1, 1_000_000],
                "annual_flow_min_t_per_year": ["ignored", "ignored"],
                "annual_flow_max_t_per_year": ["ignored", "ignored"],
                "geometry": [b"\x00", b"\x00"]}).to_parquet(source, index=False)
            for mode in ("approximation", "both", "benchmark"):
                with self.subTest(mode=mode):
                    settings = replace(self.config, input_path=source, output_dir=Path(d)/mode,
                        mode=mode, input_row=2, from_id_column="emitter_name",
                        to_id_column="sink_name", annual_flow_column="co2_t_per_yr")
                    outputs = run(settings)
                    self.assertEqual(set(outputs), {settings.output_workbook_name, "routes_costed.parquet"})
                    copied = pd.read_parquet(outputs["routes_costed.parquet"])
                    self.assertEqual(len(copied), 2)
                    self.assertEqual(copied["geometry"].tolist(), [b"\x00", b"\x00"])
                    self.assertTrue(pd.isna(copied.iloc[0]["pipeline_annual_cost_eur_per_year"]))
                    self.assertAlmostEqual(copied.iloc[1]["pipeline_annual_cost_eur_per_year"], 15463744.110506363, places=5)
                    for key in ("fixed_cost_eur_per_year", "flow_cost_eur_per_t", "annual_cost_eur_per_year", "average_cost_eur_per_t"):
                        old = "pipeline_" + key
                        new = old.replace("_eur", "_eur2024")
                        self.assertTrue(pd.isna(copied.iloc[0][new]))
                        if pd.isna(copied.iloc[1][old]):
                            self.assertTrue(pd.isna(copied.iloc[1][new]))
                        else:
                            self.assertAlmostEqual(copied.iloc[1][new], copied.iloc[1][old] * settings.conversion_factor)
                    workbook = load_workbook(outputs[settings.output_workbook_name])
                    self.assertNotIn("fit_samples", workbook.sheetnames)
                    self.assertNotIn("fit_diagnostics", workbook.sheetnames)
                    self.assertEqual(workbook["summary"].max_row, 2)
                    self.assertEqual("coefficients" in workbook.sheetnames, mode != "benchmark")
                    self.assertEqual("cost_components" in workbook.sheetnames, mode != "approximation")
                    if mode != "benchmark":
                        self.assertAlmostEqual(copied.iloc[1]["pipeline_flow_cost_eur_per_t"], 3.5576435453331564, places=10)
                    else:
                        self.assertTrue(pd.isna(copied.iloc[1]["pipeline_flow_cost_eur_per_t"]))
                    workbook.close()



if __name__ == "__main__":
    unittest.main()
