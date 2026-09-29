"""Accounting contracts and preservation of the costed source input."""
from dataclasses import replace
import hashlib
import json
from pathlib import Path
import struct
import tempfile
import unittest
import warnings
import pandas as pd
import pyarrow as pa
import pyarrow.parquet as pq
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.worksheet.table import Table
from pipeline_cost_model import PipelineCostConfig, PipelineCostModel, RouteInput
from pipeline_cost_model.costed_input import COSTED_INPUT_COLUMNS, write_costed_input
from pipeline_cost_model.exports import build_tables, write_results
from pipeline_cost_model.runner import run


class BaselineExportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.config = PipelineCostConfig(mode="both", coefficient_method="fixed_design")
        cls.route = RouteInput("001", "002", 200, 1.7, "case", 1_000_000, source_row=2)
        cls.model = PipelineCostModel(cls.config)
        cls.result = cls.model.calculate_route(cls.route)

    def test_component_groups_and_optimizer_rows_reconcile_independently(self):
        tables = build_tables(self.config, [self.result])
        c = self.result.benchmark.costs["baseline"]
        group_to_total = {"upfront_capex": "capex_total_eur",
            "annualized_capex": "annualized_capex_total_eur_per_year",
            "fixed_opex": "fixed_opex_total_eur_per_year",
            "electricity_opex": "electricity_opex_total_eur_per_year"}
        rows = tables["cost_components"]
        self.assertEqual(len({r["metric"] for r in rows}), len(rows))
        self.assertFalse(any("total" in r["metric"] for r in rows))
        for group, total in group_to_total.items():
            self.assertAlmostEqual(sum(r["value"] for r in rows if r["cost_group"] == group), c[total], places=6)
        annual = sum(r["value"] for r in rows if r["cost_group"] != "upfront_capex")
        self.assertAlmostEqual(annual, tables["summary"][0]["total_annual_cost_eur_per_year"], places=6)
        coefficients = tables["coefficients"]
        totals = [r for r in coefficients if r["row_role"] == "optimizer_total"]
        parts = [r for r in coefficients if r["row_role"] == "additive_contribution"]
        self.assertEqual((len(totals), len(parts)), (1, 3))
        for field in ("fixed_coefficient_eur_per_year", "flow_coefficient_eur_per_t"):
            self.assertAlmostEqual(sum(r[field] for r in parts), totals[0][field], places=6)
        self.assertEqual({r["basis"] for r in coefficients + rows}, {"baseline"})

    def test_optional_exports_and_no_inactive_fit_tables(self):
        with tempfile.TemporaryDirectory() as d:
            config = replace(self.config, output_dir=Path(d), write_costed_input=False)
            paths = write_results(config, [self.result])
            self.assertEqual(set(paths), {config.output_workbook_name})
            book = load_workbook(paths[config.output_workbook_name])
            self.assertEqual(book.sheetnames, ["summary", "coefficients", "cost_components", "design", "settings", "units", "read_me"])
            book.close()
            paths = write_results(replace(config, write_csv=True, write_json=True), [self.result])
            payload = json.loads(paths["pipeline_results.json"].read_text())
            self.assertEqual(payload["schema_version"], "0.5.0")
            self.assertNotIn("routes", payload)
            self.assertEqual({p.stem for p in paths.values() if p.suffix == ".csv"}, set(payload["tables"]))
            self.assertFalse(any("spatial" in r.get("basis", "") for rows in payload["tables"].values() for r in rows))

    def test_parquet_preserves_geo_metadata_nested_data_dtypes_and_saved_index(self):
        with tempfile.TemporaryDirectory() as d:
            source = Path(d)/"input.geoparquet"
            geometry = struct.pack("<BIdd", 1, 1, 4.9, 52.1)
            original_frame = pd.DataFrame({"from_id": ["001", "003", "004"], "to_id": ["002"]*3,
                "distance_km": [200.0]*3, "annual_flow_t_per_year": [1_000_000]*3,
                "geometry": [geometry]*3, "bbox": [{"xmin": 4.9}]*3, "bbox.xmin": [4.9]*3,
                "category": pd.Categorical(["A", "B", "A"]),
                "pipeline_fixed_cost_eur_per_year": [123.0]*3},
                index=pd.Index([8, 8, 2], name="saved_index"))
            table = pa.Table.from_pandas(original_frame, preserve_index=True)
            geo = json.dumps({"version": "1.0.0", "primary_column": "geometry", "columns": {
                "geometry": {"encoding": "WKB", "geometry_types": ["Point"], "crs": None}}}).encode()
            metadata = {**table.schema.metadata, b"geo": geo, b"custom_metadata": b"retained"}
            pq.write_table(table.replace_schema_metadata(metadata), source, row_group_size=1)
            config = replace(self.config, input_path=source, output_dir=Path(d)/"out", input_row=2)
            before = hashlib.sha256(source.read_bytes()).hexdigest()
            paths = run(config)
            copied = pq.read_table(paths["input_costed.geoparquet"])
            self.assertEqual(copied.schema.metadata[b"geo"], geo)
            self.assertEqual(copied.schema.metadata[b"custom_metadata"], b"retained")
            for key in original_frame.columns:
                if key not in COSTED_INPUT_COLUMNS:
                    self.assertEqual(copied.schema.field(key), table.schema.field(key))
                    self.assertEqual(copied[key].to_pylist(), table[key].to_pylist())
            frame = copied.to_pandas()
            self.assertEqual(frame.index.tolist(), [8, 8, 2])
            self.assertEqual(frame["category"].dtype, original_frame["category"].dtype)
            self.assertTrue(pd.isna(frame.iloc[0]["pipeline_fixed_cost_eur_per_year"]))
            self.assertAlmostEqual(frame.iloc[1]["pipeline_fixed_cost_eur_per_year"], 11906100.565173205, places=5)
            self.assertTrue(pd.isna(frame.iloc[2]["pipeline_fixed_cost_eur_per_year"]))
            self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest(), before)
            self.assertEqual(set(frame.columns)-set(original_frame.columns), set(COSTED_INPUT_COLUMNS)-set(original_frame.columns))

    def test_parquet_physical_row_selection_across_streaming_batch_boundary(self):
        with tempfile.TemporaryDirectory() as d:
            source = Path(d)/"large.parquet"
            original = pa.table({"sequence": pa.array(range(65539), type=pa.int32())})
            pq.write_table(original, source, row_group_size=20000)
            result = replace(self.result, route=replace(self.route, source_row=65538))
            path = Path(d)/"large_costed.parquet"
            write_costed_input(path, replace(self.config, input_path=source), [result])
            table = pq.read_table(path)
            self.assertEqual(table["sequence"].to_pylist(), original["sequence"].to_pylist())
            values = table["pipeline_cost_basis"].to_pylist()
            self.assertEqual([i for i, value in enumerate(values) if value is not None], [65536])

    def test_excel_copy_preserves_source_and_expands_full_input_table(self):
        with tempfile.TemporaryDirectory() as d:
            source = Path(d)/"input.xlsx"
            book = Workbook()
            sheet = book.active
            sheet.title = "pipeline_route_metrics"
            sheet.append(["from_id", "to_id", "distance_km", "annual_flow_t_per_year"])
            sheet.append(["001", "002", 200, 1_000_000])
            sheet.append(["003", "002", 200, 1_000_000])
            sheet["A2"].font = Font(bold=True, color="123456")
            sheet.add_table(Table(displayName="Routes", ref="A1:D3"))
            book.create_sheet("Notes")["A1"] = "=1+2"
            book.save(source)
            book.close()
            before = hashlib.sha256(source.read_bytes()).hexdigest()
            paths = run(replace(self.config, input_path=source, output_dir=Path(d)/"out", input_row=1))
            copied = load_workbook(paths["input_costed.xlsx"])
            sheet = copied["pipeline_route_metrics"]
            self.assertEqual(sheet.max_column, 17)
            self.assertEqual(sheet.tables["Routes"].ref, "A1:Q3")
            self.assertEqual(len(sheet.tables["Routes"].tableColumns), 17)
            self.assertEqual(copied["Notes"]["A1"].value, "=1+2")
            self.assertEqual(sheet["A2"].font.color.rgb, "00123456")
            self.assertTrue(sheet["A2"].font.bold)
            self.assertTrue(all(sheet.cell(3, col).value is None for col in range(5, 18)))
            copied.close()
            self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest(), before)

    def test_legacy_yaml_drops_financial_resistance_setting(self):
        with tempfile.TemporaryDirectory() as d:
            path = Path(d)/"old.yaml"
            path.write_text("spatial_fixed_opex_mode: scale_with_capex\ndetailed_workbook_name: details.xlsx\noutput_workbook_name: input_costed.xlsx\n")
            with warnings.catch_warnings(record=True) as notices:
                warnings.simplefilter("always")
                config = PipelineCostConfig.from_yaml(path)
            self.assertEqual(len(notices), 2)
            self.assertFalse(hasattr(config, "spatial_fixed_opex_mode"))
            self.assertEqual(config.output_workbook_name, "details.xlsx")
            self.assertFalse(config.write_csv or config.write_json)
            with self.assertRaisesRegex(ValueError, "Unknown cost basis"):
                self.result.approximation.predict(1_000_000, basis="spatial")

    def test_report_cannot_overwrite_source(self):
        with tempfile.TemporaryDirectory() as d:
            source = Path(d)/"input.xlsx"
            source.write_bytes(b"original")
            config = replace(self.config, input_path=source, output_dir=Path(d), output_workbook_name=source.name)
            with self.assertRaisesRegex(ValueError, "overwrite the input"):
                write_results(config, [self.result])
            self.assertEqual(source.read_bytes(), b"original")


if __name__ == "__main__":
    unittest.main()
