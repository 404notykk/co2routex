"""Regression, accounting and workflow checks. Run with unittest discovery."""
from __future__ import annotations
import contextlib
import csv
import io
import json
import math
import tempfile
import unittest
from dataclasses import replace
from pathlib import Path
from openpyxl import Workbook, load_workbook
from pipeline_cost_model import PipelineCostConfig, PipelineCostModel, RouteInput
from pipeline_cost_model.annual_costs import COMPONENT_FIELDS, fit_affine
from pipeline_cost_model.exports import write_results
from pipeline_cost_model.inputs import read_routes
from pipeline_cost_model.runner import run


class EngineeringTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.config = PipelineCostConfig(mode="benchmark")
        cls.model = PipelineCostModel(cls.config)
        cls.route = RouteInput("A", "B", 200, 1, "known_case", 1e6)
        cls.point = cls.model.calculate_fixed_flow(cls.route, 1e6)

    def test_reference_case_and_transported_flow_units(self):
        p = self.point
        self.assertEqual(p.operating_flow_t_per_h, 125)
        self.assertAlmostEqual(p.massflow_kg_per_s, 34.72222222222222)
        self.assertAlmostEqual(p.costs["baseline"]["capex_total_eur"], 96637270.51674786, places=5)
        self.assertAlmostEqual(p.costs["baseline"]["levelized_cost_eur_per_t"], 15.463744110506363, places=9)
        self.assertEqual(p.design["number_of_booster_stations"], 3)
        self.assertGreater(p.design["outer_diameter_m"], p.design["inner_diameter_m"])

    def test_selected_pipeline_components_reconcile(self):
        c = self.point.costs["baseline"]
        self.assertAlmostEqual(c["capex_material_eur"] + c["capex_labour_eur"] + c["capex_row_eur"] + c["capex_miscellaneous_eur"], c["capex_pipeline_eur"], places=5)
        self.assertAlmostEqual(c["capex_row_eur"], 19866460, places=5)
        self.assertAlmostEqual(c["fixed_opex_pipeline_eur_per_year"], .015 * c["capex_pipeline_eur"], places=6)

    def test_booster_energy_reporting_regression(self):
        p = self.point
        self.assertAlmostEqual(p.energy["specific_total_mwh_per_t"], .059294059088885936, places=12)
        self.assertAlmostEqual(p.energy["total_mwh_per_year"] * 60, p.costs["baseline"]["electricity_opex_total_eur_per_year"], places=6)
        self.assertAlmostEqual(p.energy["total_power_mw"] * 8000, p.energy["total_mwh_per_year"], places=7)

    def test_corrected_source_pipeline_opex_field(self):
        raw = next(iter(self.model._source_cache.values()))
        self.assertEqual(raw["cost_pipeline"]["opex_fix_abs"], raw["configuration"]["opex_pipe"])
        self.assertAlmostEqual(raw["configuration"]["opex_fix"], raw["configuration"]["opex_pipe"] + raw["configuration"]["opex_fix_compression"])

    def test_offshore_zero_row_and_boosters(self):
        p = self.model.calculate_fixed_flow(replace(self.route, terrain="Offshore"), 1e6)
        self.assertEqual(p.costs["baseline"]["capex_row_eur"], 0)
        self.assertEqual(p.design["number_of_booster_stations"], 0)
        self.assertEqual(p.energy["boosters_mwh_per_year"], 0)
        self.assertAlmostEqual(p.costs["baseline"]["levelized_cost_eur_per_t"], 13.879338421002359, places=9)

    def test_resistance_is_metadata_and_costs_remain_baseline_only(self):
        for resistance in (0.25, 2, 30):
            p = self.model.calculate_fixed_flow(replace(self.route, average_route_resistance=resistance), 1e6)
            self.assertEqual(set(p.costs), {"baseline"})
            self.assertEqual(p.costs, self.point.costs)
            self.assertEqual(p.design, self.point.design)
            self.assertEqual(p.energy, self.point.energy)
        self.assertEqual(len(self.model._source_cache), 2 if any(key[2] == "Offshore" for key in self.model._source_cache) else 1)

    def test_zero_electricity_price_preserves_energy(self):
        p = PipelineCostModel(replace(self.config, electricity_price_eur_per_mwh=0)).calculate_fixed_flow(self.route, 1e6)
        self.assertGreater(p.energy["total_mwh_per_year"], 0)
        self.assertEqual(p.costs["baseline"]["electricity_opex_total_eur_per_year"], 0)
        c = p.costs["baseline"]
        self.assertEqual(c["total_annual_cost_eur_per_year"], c["annualized_capex_total_eur_per_year"] + c["fixed_opex_total_eur_per_year"])

    def test_benchmark_ignores_fit_bounds_and_has_no_old_capacity_floor(self):
        config = replace(self.config, annual_flow_min_t_per_year=0, annual_flow_max_t_per_year=0)
        p = PipelineCostModel(config).calculate_route(replace(self.route, annual_flow_t_per_year=100000))
        self.assertIsNone(p.approximation)
        self.assertEqual(p.benchmark.operating_flow_t_per_h, 12.5)


class ApproximationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.config = PipelineCostConfig(mode="both", annual_flow_min_t_per_year=500000, annual_flow_max_t_per_year=2000000)
        cls.route = RouteInput("A", "B", 200, 1.7, "affine", 1e6)
        cls.model = PipelineCostModel(cls.config)
        cls.result = cls.model.calculate_route(cls.route)

    def test_endpoint_fit_uses_total_annual_cost_and_separate_components(self):
        a = self.result.approximation
        for basis in ("baseline",):
            for role, point in a.samples:
                if role == "fit":
                    self.assertAlmostEqual(a.predict(point.annual_flow_t_per_year, basis=basis), point.costs[basis]["total_annual_cost_eur_per_year"], places=5)
            coefs = a.coefficients[basis]
            for attr in ("intercept_eur_per_year", "slope_eur_per_t"):
                self.assertAlmostEqual(getattr(coefs["total"], attr), getattr(coefs["annualized_capex"], attr) + getattr(coefs["opex"], attr), places=6)
                self.assertAlmostEqual(getattr(coefs["opex"], attr), getattr(coefs["fixed_opex"], attr) + getattr(coefs["electricity_opex"], attr), places=6)

    def test_direct_benchmark_is_not_replaced_by_fit(self):
        r = self.result
        direct = r.benchmark.costs["baseline"]["total_annual_cost_eur_per_year"]
        self.assertAlmostEqual(direct, 15463744.110506363, places=5)
        self.assertGreater(abs(r.approximation.predict(1e6) - direct), 100000)
        self.assertIn("benchmark_check", [role for role, _ in r.approximation.samples])

    def test_validation_and_flow_range_constraints(self):
        a = self.result.approximation
        fit_flows = {p.annual_flow_t_per_year for role, p in a.samples if role == "fit"}
        validation = {p.annual_flow_t_per_year for role, p in a.samples if role != "fit"}
        self.assertFalse(fit_flows & validation)
        self.assertGreater(a.diagnostics["baseline"]["total"]["validation_max_abs_relative_error_percent"], 0)
        self.assertEqual(a.predict(0, built=False), 0)
        with self.assertRaises(ValueError):
            a.predict(0)
        with self.assertRaises(ValueError):
            a.predict(3000000)

    def test_least_squares_and_error_warning(self):
        config = replace(self.config, fit_method="least_squares", fit_samples=5, validation_samples=4, fit_error_warning_percent=.01)
        r = PipelineCostModel(config).calculate_route(self.route)
        a = r.approximation
        fitted = [p for role, p in a.samples if role == "fit"]
        self.assertEqual(len(fitted), 5)
        # Least-squares residuals sum to zero when an intercept is included.
        residual_sum = sum(a.predict(p.annual_flow_t_per_year) - p.costs["baseline"]["total_annual_cost_eur_per_year"] for p in fitted)
        self.assertAlmostEqual(residual_sum, 0, places=5)
        self.assertTrue(any("exceeds" in warning for warning in a.warnings))

    def test_hourly_legacy_bounds_convert_to_annual_mass(self):
        config = PipelineCostConfig(mode="approximation", capacity_min_t_per_h=18, capacity_max_t_per_h=4050)
        lo, hi, _ = config.flow_bounds(RouteInput("A", "B", 10))
        self.assertEqual((lo, hi), (144000, 32400000))


class WorkflowTests(unittest.TestCase):
    def test_explicit_capturable_column_and_scenario_identity(self):
        with tempfile.TemporaryDirectory() as d:
            source = Path(d) / "input.csv"
            source.write_text("scenario_id,from_id,to_id,distance_km,co2_t_annual,co2_capturable_t_annual\ns1,001,002,200,1000000,900000\ns2,001,002,200,2000000,1800000\n")
            config = PipelineCostConfig(input_path=source, mode="benchmark", annual_flow_column="co2_capturable_t_annual")
            rows = read_routes(config)
            self.assertEqual([r.annual_flow_t_per_year for r in rows], [900000, 1800000])
            self.assertEqual(rows[0].from_id, "001")
            with self.assertRaisesRegex(ValueError, "benchmark requires"):
                read_routes(replace(config, annual_flow_column="annual_flow_t_per_year"))
            source.write_text(source.read_text().replace("s2,", "s1,"))
            with self.assertRaisesRegex(ValueError, "Duplicate scenario"):
                read_routes(config)

    def test_excel_preservation_filtered_rows_and_mode_refresh(self):
        with tempfile.TemporaryDirectory() as d:
            source = Path(d) / "input.xlsx"
            w = Workbook()
            sheet = w.active
            sheet.title = "pipeline_route_metrics"
            sheet.append(["scenario_id", "from_id", "to_id", "distance_km", "annual_flow_t_per_year", "status"])
            sheet.append(["first", "A", "B", 200, 1000000, "ok"])
            sheet.append(["skip", "A", "B", 200, 1000000, "failed"])
            sheet.append(["second", "A", "B", 200, 1500000, "ok"])
            w.create_sheet("Preserve")["A1"] = "=1+2"
            w.save(source)
            w.close()
            config = PipelineCostConfig(input_path=source, output_dir=Path(d)/"results", mode="both",
                annual_flow_min_t_per_year=500000, annual_flow_max_t_per_year=2000000, validation_samples=2)
            outputs = run(config)
            w = load_workbook(outputs["input_costed.xlsx"])
            self.assertEqual(w["Preserve"]["A1"].value, "=1+2")
            sheet = w["pipeline_route_metrics"]
            headers = {c.value: c.column for c in sheet[1]}
            key = "pipeline_annual_cost_eur_per_year"
            self.assertIsNone(sheet.cell(3, headers[key]).value)
            self.assertNotEqual(sheet.cell(2, headers[key]).value, sheet.cell(4, headers[key]).value)
            w.close()
            self.assertEqual(set(outputs), {config.output_workbook_name, "input_costed.xlsx"})
            report = load_workbook(outputs[config.output_workbook_name])
            self.assertEqual(report["summary"].max_row, 3)
            self.assertIn("coefficients", report.sheetnames)
            report.close()
            # A benchmark-only rerun clears the old coefficient outputs.
            outputs = run(replace(config, mode="benchmark"))
            report = load_workbook(outputs[config.output_workbook_name])
            self.assertNotIn("coefficients", report.sheetnames)
            report.close()
            w = load_workbook(outputs["input_costed.xlsx"])
            sheet = w["pipeline_route_metrics"]
            headers = {c.value: c.column for c in sheet[1]}
            self.assertIsNone(sheet.cell(2, headers["pipeline_fixed_cost_eur_per_year"]).value)
            w.close()

    def test_approximation_only_needs_no_benchmark_flow(self):
        with tempfile.TemporaryDirectory() as d:
            source = Path(d)/"input.csv"
            source.write_text("from_id,to_id,distance_km\nA,B,200\n")
            config = PipelineCostConfig(input_path=source, output_dir=Path(d)/"out", mode="approximation", write_xlsx=False, write_json=True,
                annual_flow_min_t_per_year=500000, annual_flow_max_t_per_year=2000000, validation_samples=2)
            outputs = run(config)
            self.assertFalse(any(p.suffix == ".xlsx" for p in outputs.values()))
            payload = json.loads(outputs["pipeline_results.json"].read_text())
            self.assertNotIn("cost_components", payload["tables"])
            self.assertNotIn("total_annual_cost_eur_per_year", payload["tables"]["summary"][0])
            self.assertIn("coefficients", payload["tables"])


if __name__ == "__main__":
    unittest.main()
