"""Batch failures are visible results; they never become zero-cost routes."""
from dataclasses import replace
import hashlib
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
import pandas as pd

from pipeline_cost_model import PipelineCostConfig, PipelineCostModel, RouteInput
from pipeline_cost_model.costed_input import COSTED_INPUT_COLUMNS
from pipeline_cost_model.runner import run


class BatchFailureTests(unittest.TestCase):
    annual = 90.4028 * 3.6 * 8000

    def frame(self):
        return pd.DataFrame({
            "from_id": ["before", "Uniper Maasvlakte", "after", "filtered"],
            "to_id": ["Aramis"] * 4,
            "distance_km": [200.0, 0.241421, 103.8065150626205, 0.0],
            "annual_flow_t_per_year": [1e6, self.annual, 3278880.0, 0.0],
            "status": ["ok"] * 4, "mode": ["pipeline"] * 3 + ["truck"],
            # Results from an earlier run must not survive a current failure.
            "pipeline_fixed_cost_eur_per_year": [123.0] * 4,
            "pipeline_fixed_cost_eur2024_per_year": [456.0] * 4,
            "pipeline_flow_cost_eur2024_per_t": [7.0] * 4,
            "pipeline_annual_cost_eur2024_per_year": [789.0] * 4,
            "pipeline_average_cost_eur2024_per_t": [8.0] * 4,
            "pipeline_cost_status": ["failed"] * 4,
            "pipeline_cost_error": ["stale message"] * 4,
        }, index=pd.Index([8, 8, 2, 0], name="saved_index"))

    def config(self, source, directory, **kwargs):
        return PipelineCostConfig(input_path=source, output_dir=directory,
                                  coefficient_method="fixed_design", **kwargs)

    def test_real_failure_continues_and_exports_only_successful_costs_in_each_mode(self):
        for mode in ("approximation", "benchmark", "both"):
            with self.subTest(mode=mode), tempfile.TemporaryDirectory() as directory:
                root = Path(directory)
                source = root / "routes.parquet"
                self.frame().to_parquet(source)
                before = hashlib.sha256(source.read_bytes()).hexdigest()
                config = self.config(source, root / "out", mode=mode)
                with self.assertLogs("pipeline_cost_model.runner", level="INFO") as messages:
                    outputs = run(config)
                self.assertEqual(set(outputs), {"pipeline_cost_results.xlsx", "routes_costed.parquet"})
                self.assertTrue(any("2 successful, 1 failed, 3 routes attempted" in line for line in messages.output))
                copied = pd.read_parquet(outputs["routes_costed.parquet"])
                self.assertEqual(copied.index.tolist(), [8, 8, 2, 0])
                self.assertEqual(copied.iloc[:3]["pipeline_cost_status"].tolist(), ["ok", "failed", "ok"])
                self.assertTrue(pd.isna(copied.iloc[3]["pipeline_cost_status"]))
                self.assertIn("No different flow has been substituted", copied.iloc[1]["pipeline_cost_error"])
                for column in COSTED_INPUT_COLUMNS:
                    if column in {"pipeline_cost_status", "pipeline_cost_error"}:
                        continue
                    self.assertTrue(pd.isna(copied.iloc[1][column]), column)
                    self.assertTrue(pd.isna(copied.iloc[3][column]), column)
                self.assertTrue(pd.isna(copied.iloc[0]["pipeline_cost_error"]))
                self.assertTrue(pd.isna(copied.iloc[2]["pipeline_cost_error"]))
                self.assertAlmostEqual(copied.iloc[0]["pipeline_annual_cost_eur_per_year"],
                                       15463744.110506363, places=6)
                # This route is calculated after the failure, with a different
                # distance and flow, so it cannot reuse the earlier cached case.
                self.assertAlmostEqual(copied.iloc[2]["pipeline_annual_cost_eur_per_year"],
                                       22085822.76767799, places=6)
                book = load_workbook(outputs["pipeline_cost_results.xlsx"], data_only=True)
                rows = list(book["summary"].values)
                summary = [dict(zip(rows[0], row)) for row in rows[1:]]
                self.assertEqual([row["input_data_row"] for row in summary], [1, 2, 3])
                self.assertEqual([row["cost_status"] for row in summary], ["ok", "failed", "ok"])
                self.assertIsNone(summary[1]["total_annual_cost_eur_per_year"])
                self.assertIsNone(summary[1]["total_annual_cost_eur2024_per_year"])
                self.assertEqual(summary[1]["input_annual_flow_t_per_year"], self.annual)
                self.assertIn("engineering calculation failed", summary[1]["cost_error"])
                for sheet in ("coefficients", "cost_components", "design"):
                    if sheet in book.sheetnames:
                        data = list(book[sheet].values)
                        scenario_column = data[0].index("scenario_id")
                        self.assertEqual({row[scenario_column] for row in data[1:]}, {"row_2", "row_4"})
                settings = dict(list(book["settings"].values)[1:])
                self.assertEqual(settings["routes_failed"], "1")
                self.assertEqual(settings["routes_successful"], "2")
                book.close()
                self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest(), before)

    def test_selected_failed_row_and_all_failed_batch_still_produce_outputs(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "routes.parquet"
            self.frame().to_parquet(source)
            outputs = run(self.config(source, root / "out", mode="approximation",
                                      input_row=2, write_json=True, write_csv=True))
            payload = json.loads(outputs["pipeline_results.json"].read_text())
            tables = payload["tables"]
            self.assertEqual(set(tables), {"summary", "settings", "units"})
            self.assertEqual(tables["summary"][0]["cost_status"], "failed")
            self.assertIsNone(tables["summary"][0]["total_annual_cost_eur_per_year"])
            self.assertIsNone(tables["summary"][0]["total_annual_cost_eur2024_per_year"])
            self.assertNotIn("coefficients.csv", outputs)
            copied = pd.read_parquet(outputs["routes_costed.parquet"])
            self.assertEqual(copied["pipeline_cost_status"].dropna().tolist(), ["failed"])
            self.assertTrue(copied["pipeline_annual_cost_eur_per_year"].isna().all())
            for column in COSTED_INPUT_COLUMNS:
                if "eur2024" in column:
                    self.assertTrue(copied[column].isna().all(), column)
            book = load_workbook(outputs["pipeline_cost_results.xlsx"])
            self.assertEqual(book.sheetnames, ["summary", "settings", "units", "read_me"])
            book.close()

    def test_failure_status_survives_csv_and_excel_costed_copies(self):
        for extension in ("csv", "xlsx"):
            with self.subTest(extension=extension), tempfile.TemporaryDirectory() as directory:
                root = Path(directory)
                source = root / f"routes.{extension}"
                frame = self.frame()
                if extension == "csv":
                    frame.to_csv(source, index=False)
                else:
                    book = Workbook()
                    sheet = book.active
                    sheet.title = "pipeline_route_metrics"
                    sheet.append(frame.columns.tolist())
                    for row in frame.itertuples(index=False, name=None):
                        sheet.append(row)
                    book.save(source)
                    book.close()
                outputs = run(self.config(source, root / "out", input_row=2,
                                          mode="approximation", write_xlsx=False))
                path = outputs[f"routes_costed.{extension}"]
                if extension == "csv":
                    copied = pd.read_csv(path)
                else:
                    copied = pd.read_excel(path, sheet_name="pipeline_route_metrics")
                self.assertEqual(copied.loc[1, "pipeline_cost_status"], "failed")
                self.assertIn("No different flow", copied.loc[1, "pipeline_cost_error"])
                self.assertTrue(copied["pipeline_annual_cost_eur_per_year"].isna().all())
                for column in COSTED_INPUT_COLUMNS:
                    if "eur2024" in column:
                        self.assertTrue(copied[column].isna().all(), column)
                self.assertTrue(pd.isna(copied.loc[0, "pipeline_cost_error"]))

    def test_cli_reports_failures_but_finishes_after_writing_results(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            self.frame().to_csv(root / "routes.csv", index=False)
            config = root / "config.yaml"
            config.write_text("input_path: routes.csv\noutput_dir: out\nwrite_xlsx: false\n")
            cli = Path(__file__).resolve().parents[1] / "run_pipeline_cost_model.py"
            completed = subprocess.run([sys.executable, str(cli), "--config", str(config),
                                        "--mode", "approximation", "--coefficient-method", "fixed_design"],
                                       capture_output=True, text=True, timeout=30)
            self.assertEqual(completed.returncode, 0, completed.stderr)
            self.assertIn("2 successful, 1 failed, 3 routes attempted", completed.stderr)
            self.assertIn("Results written", completed.stdout)
            self.assertTrue((root / "out" / "routes_costed.csv").is_file())

    def test_interruption_and_resource_errors_are_not_swallowed(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "routes.csv"
            self.frame().to_csv(source, index=False)
            config = self.config(source, root / "out", mode="approximation")
            for error in (KeyboardInterrupt, SystemExit, MemoryError):
                with self.subTest(error=error), patch.object(PipelineCostModel, "calculate_route", side_effect=error):
                    with self.assertRaises(error):
                        run(config)
            self.assertFalse(config.output_dir.exists())

    def test_input_and_export_errors_remain_fatal(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "routes.csv"
            source.write_text("unrelated\nvalue\n")
            config = self.config(source, root / "out", mode="approximation")
            with self.assertRaisesRegex(ValueError, "Missing input columns"):
                run(config)
            self.frame().to_csv(source, index=False)
            with patch("pipeline_cost_model.runner.write_results", side_effect=OSError("disk full")):
                with self.assertRaisesRegex(OSError, "disk full"):
                    run(config)

    def test_direct_calculation_keeps_original_short_route_error(self):
        model = PipelineCostModel(PipelineCostConfig(coefficient_method="fixed_design"))
        route = RouteInput("Uniper Maasvlakte", "Aramis", 0.241421, 1, "row_6", self.annual)
        with self.assertRaisesRegex(ValueError, "No different flow has been substituted"):
            model.calculate_route(route)


if __name__ == "__main__":
    unittest.main()
