"""Write MILP-facing coefficients and separate endpoint details."""

from __future__ import annotations

import os
import shutil
from dataclasses import asdict
from pathlib import Path
from uuid import uuid4

import pandas as pd
from openpyxl import load_workbook

from .config import PipelineCostConfig
from .models import RouteCostResult


COST_COLUMNS = [
    "capacity_min_t_per_h",
    "capacity_max_t_per_h",
    "capex_baseline_slope_eur_per_t_per_h",
    "capex_baseline_intercept_eur",
    "capex_spatial_slope_eur_per_t_per_h",
    "capex_spatial_intercept_eur",
    "currency",
    "cost_reference_year",
    "terrain",
    "capacity_range_basis",
]


def write_results(config: PipelineCostConfig, results: list[RouteCostResult]) -> tuple[Path, Path]:
    config.output_dir.mkdir(parents=True, exist_ok=True)
    main_path = config.output_dir / config.output_workbook_name
    details_path = config.output_dir / config.detailed_workbook_name
    if main_path.resolve() == config.workbook_path.resolve():
        raise ValueError("Output workbook must not overwrite the input workbook")
    if main_path.resolve() == details_path.resolve():
        raise ValueError("Main and detailed workbook paths must differ")

    temporary_main = _temporary_path(main_path)
    temporary_details = _temporary_path(details_path)
    try:
        shutil.copy2(config.workbook_path, temporary_main)
        _append_coefficients(temporary_main, config, results)
        _write_details(temporary_details, config, results)
        os.replace(temporary_main, main_path)
        os.replace(temporary_details, details_path)
    finally:
        temporary_main.unlink(missing_ok=True)
        temporary_details.unlink(missing_ok=True)
    return main_path, details_path


def _append_coefficients(path: Path, config: PipelineCostConfig, results: list[RouteCostResult]) -> None:
    workbook = load_workbook(path)
    if config.route_metrics_sheet not in workbook.sheetnames:
        workbook.close()
        raise ValueError(f"Workbook does not contain {config.route_metrics_sheet!r}")
    sheet = workbook[config.route_metrics_sheet]
    headers = {str(sheet.cell(1, col).value).strip(): col for col in range(1, sheet.max_column + 1)}
    for required in ("from_id", "to_id"):
        if required not in headers:
            workbook.close()
            raise ValueError(f"{config.route_metrics_sheet!r} lacks {required!r}")

    next_column = sheet.max_column + 1
    for name in COST_COLUMNS:
        if name not in headers:
            headers[name] = next_column
            sheet.cell(1, next_column, name)
            next_column += 1

    rows = {}
    for row in range(2, sheet.max_row + 1):
        key = (_clean_id(sheet.cell(row, headers["from_id"]).value), _clean_id(sheet.cell(row, headers["to_id"]).value))
        rows[key] = row

    for result in results:
        key = (result.route.from_id, result.route.to_id)
        if key not in rows:
            workbook.close()
            raise ValueError(f"Cannot find route {key[0]} -> {key[1]} in output sheet")
        values = {
            "capacity_min_t_per_h": result.minimum.capacity_t_per_h,
            "capacity_max_t_per_h": result.maximum.capacity_t_per_h,
            "capex_baseline_slope_eur_per_t_per_h": result.capex_baseline_slope_eur_per_t_per_h,
            "capex_baseline_intercept_eur": result.capex_baseline_intercept_eur,
            "capex_spatial_slope_eur_per_t_per_h": result.capex_spatial_slope_eur_per_t_per_h,
            "capex_spatial_intercept_eur": result.capex_spatial_intercept_eur,
            "currency": config.currency,
            "cost_reference_year": config.cost_reference_year,
            "terrain": config.terrain,
            "capacity_range_basis": config.capacity_range_basis,
        }
        for name, value in values.items():
            cell = sheet.cell(rows[key], headers[name], value)
            if isinstance(value, float):
                cell.number_format = "0.000000"
    workbook.save(path)
    workbook.close()


def _write_details(path: Path, config: PipelineCostConfig, results: list[RouteCostResult]) -> None:
    endpoint_rows = []
    coefficient_rows = []
    for result in results:
        coefficient_rows.append({
            "from_id": result.route.from_id,
            "to_id": result.route.to_id,
            "distance_km": result.route.distance_km,
            "average_route_resistance": result.route.average_route_resistance,
            "capex_baseline_slope_eur_per_t_per_h": result.capex_baseline_slope_eur_per_t_per_h,
            "capex_baseline_intercept_eur": result.capex_baseline_intercept_eur,
            "capex_spatial_slope_eur_per_t_per_h": result.capex_spatial_slope_eur_per_t_per_h,
            "capex_spatial_intercept_eur": result.capex_spatial_intercept_eur,
        })
        for label, endpoint in (("minimum", result.minimum), ("maximum", result.maximum)):
            endpoint_rows.append({
                "from_id": result.route.from_id,
                "to_id": result.route.to_id,
                "endpoint": label,
                "distance_km": result.route.distance_km,
                "average_route_resistance": result.route.average_route_resistance,
                **asdict(endpoint),
            })

    settings = [{"parameter": key, "value": str(value)} for key, value in asdict(config).items()]
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        pd.DataFrame(coefficient_rows).to_excel(writer, sheet_name="route_coefficients", index=False)
        pd.DataFrame(endpoint_rows).to_excel(writer, sheet_name="endpoint_calculations", index=False)
        pd.DataFrame(settings).to_excel(writer, sheet_name="configuration", index=False)


def _clean_id(value) -> str:
    text = str(value).strip()
    if text.endswith(".0") and text[:-2].isdigit():
        return text[:-2]
    return text


def _temporary_path(target: Path) -> Path:
    return target.with_name(f".{target.stem}.{uuid4().hex}{target.suffix}")
