"""One baseline report, a costed input copy, and optional machine-readable tables."""
from __future__ import annotations
import csv
import json
import os
import tempfile
from dataclasses import asdict
from datetime import datetime, timezone
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from .annual_costs import FITTED_COMPONENTS
from .models import Approximation, FixedDesignApproximation
from .costed_input import costed_input_name, write_costed_input
from .currency import add_eur2024_columns, conversion_metadata, to_eur2024

VERSION = "0.5.0"
EXPORTED_COMPONENTS = (*FITTED_COMPONENTS, "total")
# Entries are additive within their cost_group. Totals live in summary.
COST_COMPONENTS = (
    ("upfront_capex", "pipeline", "Material", "capex_material_eur"),
    ("upfront_capex", "pipeline", "Labour", "capex_labour_eur"),
    ("upfront_capex", "pipeline", "Right of way", "capex_row_eur"),
    ("upfront_capex", "pipeline", "Miscellaneous", "capex_miscellaneous_eur"),
    ("upfront_capex", "initial_compression", "Equipment", "capex_initial_compression_eur"),
    ("upfront_capex", "boosters", "Equipment", "capex_boosters_eur"),
    ("annualized_capex", "pipeline", "Capital recovery", "annualized_capex_pipeline_eur_per_year"),
    ("annualized_capex", "initial_compression", "Capital recovery", "annualized_capex_initial_compression_eur_per_year"),
    ("annualized_capex", "boosters", "Capital recovery", "annualized_capex_boosters_eur_per_year"),
    ("fixed_opex", "pipeline", "Fixed operation and maintenance", "fixed_opex_pipeline_eur_per_year"),
    ("fixed_opex", "initial_compression", "Fixed operation and maintenance", "fixed_opex_initial_compression_eur_per_year"),
    ("fixed_opex", "boosters", "Fixed operation and maintenance", "fixed_opex_boosters_eur_per_year"),
    ("electricity_opex", "initial_compression", "Electricity", "electricity_opex_initial_compression_eur_per_year"),
    ("electricity_opex", "boosters", "Electricity", "electricity_opex_boosters_eur_per_year"),
)
SUMMARY_COST_FIELDS = (
    "capex_total_eur", "annualized_capex_total_eur_per_year", "fixed_opex_total_eur_per_year",
    "electricity_opex_total_eur_per_year", "total_annual_cost_eur_per_year", "levelized_cost_eur_per_t",
)


def route_metadata(route):
    return {"scenario_id": route.scenario_id, "from_id": route.from_id, "to_id": route.to_id,
            "input_data_row": route.source_row - 1 if route.source_row is not None else None,
            "distance_km": route.distance_km, "basis": "baseline"}


def reference_point(result):
    if result.benchmark is not None:
        return result.benchmark
    if isinstance(result.approximation, FixedDesignApproximation):
        return result.approximation.reference
    return None


def unit_for(metric):
    if metric == "value":
        return "See the row's unit or parameter"
    if metric == "value_eur2024":
        return "See the row's unit_eur2024"
    for suffix, unit in (
        ("eur2024_per_mwh", "EUR_2024/MWh"), ("eur2024_per_year", "EUR_2024/year"),
        ("eur2024_per_t", "EUR_2024/tCO2"), ("_eur2024", "EUR_2024"),
        ("eur_per_mwh", "EUR_2021/MWh"), ("eur_per_year", "EUR_2021/year"),
        ("eur_per_t", "EUR_2021/tCO2"), ("_eur", "EUR_2021"),
        ("mwh_per_year", "MWh/year"), ("mwh_per_t", "MWh/tCO2"), ("_mw", "MW"),
        ("t_per_year", "tCO2/year"), ("t_per_h", "tCO2/h"), ("kg_per_s", "kgCO2/s"),
        ("_percent", "%"), ("_km", "km"), ("_bar", "bar"), ("_m", "m"), ("_years", "year"),
    ):
        if metric.endswith(suffix):
            return unit
    if metric == "operating_hours_per_year":
        return "h/year"
    if metric.endswith("_fraction") or metric == "discount_rate":
        return "fraction"
    return "dimensionless/text"


def build_tables(config, results):
    """Build each view once. Omit inactive views rather than writing empty files."""
    factor = config.conversion_factor
    tables = {name: [] for name in ("summary", "coefficients", "cost_components", "design",
                                   "fit_samples", "fit_diagnostics")}
    for result in results:
        meta = route_metadata(result.route)
        a = result.approximation
        point = reference_point(result)
        main = {**meta, "run_mode": config.mode, "cost_model_version": VERSION,
                "coefficient_method": config.coefficient_method if config.mode != "benchmark" else None,
                "cost_status": result.cost_status, "cost_error": result.error_message,
                "input_annual_flow_t_per_year": result.route.annual_flow_t_per_year}
        if result.cost_status == "failed":
            main.update({"reference_annual_flow_t_per_year": None, **dict.fromkeys(SUMMARY_COST_FIELDS)})
            tables["summary"].append(main)
            continue
        if point is not None:
            main["reference_annual_flow_t_per_year"] = point.annual_flow_t_per_year
            main.update({key: point.costs["baseline"][key] for key in SUMMARY_COST_FIELDS})
            tables["design"].append({**meta, "reference_annual_flow_t_per_year": point.annual_flow_t_per_year,
                "operating_flow_t_per_h": point.operating_flow_t_per_h, "massflow_kg_per_s": point.massflow_kg_per_s,
                "terrain": point.terrain, "routing_resistance_metadata": result.route.average_route_resistance,
                **point.design, **point.energy, **point.assumptions})
        if a is not None:
            for component in EXPORTED_COMPONENTS:
                coefficient = a.coefficients["baseline"][component]
                row = {**meta, "component": component,
                    "row_role": "optimizer_total" if component == "total" else "additive_contribution",
                    "fixed_coefficient_eur_per_year": coefficient.intercept_eur_per_year,
                    "flow_coefficient_eur_per_t": coefficient.slope_eur_per_t,
                    "coefficient_method": a.coefficient_method}
                if isinstance(a, FixedDesignApproximation):
                    row["reference_annual_flow_t_per_year"] = a.reference_annual_flow_t_per_year
                else:
                    row.update({"min_annual_flow_t_per_year": a.annual_flow_min_t_per_year,
                                "max_annual_flow_t_per_year": a.annual_flow_max_t_per_year})
                tables["coefficients"].append(row)
            if isinstance(a, FixedDesignApproximation):
                main["reference_reconciliation_error_eur_per_year"] = (
                    a.predict(point.annual_flow_t_per_year) - point.costs["baseline"]["total_annual_cost_eur_per_year"])
            else:
                main.update({"min_annual_flow_t_per_year": a.annual_flow_min_t_per_year,
                             "max_annual_flow_t_per_year": a.annual_flow_max_t_per_year,
                             "fit_warnings": "; ".join(a.warnings)})
                for component in EXPORTED_COMPONENTS:
                    tables["fit_diagnostics"].append({**meta, "component": component,
                        **a.diagnostics["baseline"][component]})
                for role, sample in a.samples:
                    actual = sample.costs["baseline"]["total_annual_cost_eur_per_year"]
                    prediction = a.predict(sample.annual_flow_t_per_year)
                    tables["fit_samples"].append({**meta, "sample_role": role,
                        "annual_flow_t_per_year": sample.annual_flow_t_per_year,
                        "inner_diameter_m": sample.design["inner_diameter_m"],
                        "number_of_booster_stations": sample.design["number_of_booster_stations"],
                        "actual_total_annual_cost_eur_per_year": actual,
                        "predicted_total_annual_cost_eur_per_year": prediction,
                        "error_percent": None if actual == 0 else 100 * (prediction - actual) / actual})
                if point is not None:
                    q = point.annual_flow_t_per_year
                    inside = a.annual_flow_min_t_per_year <= q <= a.annual_flow_max_t_per_year
                    main["reference_inside_fit_range"] = inside
                    if inside:
                        main["fitted_reference_cost_eur_per_year"] = a.predict(q)
        if result.benchmark is not None:
            for group, asset, component, metric in COST_COMPONENTS:
                tables["cost_components"].append({**meta, "cost_group": group, "asset": asset,
                    "component": component, "metric": metric,
                    "value": result.benchmark.costs["baseline"][metric], "unit": unit_for(metric),
                    "value_eur2024": to_eur2024(result.benchmark.costs["baseline"][metric], factor),
                    "unit_eur2024": unit_for(metric.replace("_eur", "_eur2024", 1))})
        tables["summary"].append(main)
    tables = {name: [add_eur2024_columns(row, factor) for row in rows]
              for name, rows in tables.items() if rows}
    settings = asdict(config)
    settings.update(conversion_metadata(config))
    if not config.uses_flow_range:
        for name in ("annual_flow_min_t_per_year", "annual_flow_max_t_per_year", "capacity_min_t_per_h",
                     "capacity_max_t_per_h", "capacity_range_basis", "fit_method", "fit_samples",
                     "validation_samples", "fit_error_warning_percent"):
            settings.pop(name, None)
    tables["settings"] = [{"parameter": key, "value": str(value)} for key, value in settings.items()]
    tables["settings"] += [
        {"parameter": "cost_basis", "value": "baseline"},
        {"parameter": "resistance_use", "value": "Routing only. Physical routed distance is used without a financial multiplier."},
        {"parameter": "coefficient_equation", "value": "annual_cost = F * built + v * annual_flow"},
        {"parameter": "coefficient_meaning", "value": "F = annualized CAPEX + fixed OPEX; v = reference electricity OPEX / reference flow" if config.coefficient_method == "fixed_design" else "Affine coefficients of annual costs across re-sized designs"},
        {"parameter": "operating_assumption", "value": "Fixed design and constant electricity per tonne; 0 <= flow <= reference flow * built; off-reference hydraulics not evaluated" if config.coefficient_method == "fixed_design" else "Design varies with flow; use the exported fitted flow bounds"},
        {"parameter": "gross_emissions_conversion", "value": "None. Input flow is transported CO2 in tonnes/year."},
        {"parameter": "route_error_handling", "value": "Continue after route calculation failures; retain status and error with blank costs. Configuration, input-reading and output-writing failures stop the run."},
        {"parameter": "routes_attempted", "value": str(len(results))},
        {"parameter": "routes_successful", "value": str(sum(r.cost_status == "ok" for r in results))},
        {"parameter": "routes_failed", "value": str(sum(r.cost_status == "failed" for r in results))},
    ]
    metrics = sorted({key for rows in tables.values() for row in rows for key in row})
    tables["units"] = [{"field": key, "unit": unit_for(key)} for key in metrics]
    return tables


def report_notes():
    return [
        {"topic": "Summary", "meaning": "Annualized capital, fixed OPEX and electricity OPEX sum to total annual cost. Upfront CAPEX is a separate investment amount."},
        {"topic": "Coefficients", "meaning": "Use the optimizer_total row for each connection. The other three rows explain its composition. Do not add the total to its contributions."},
        {"topic": "Cost components", "meaning": "Each row is additive within its cost_group. Sum annualized_capex, fixed_opex and electricity_opex to obtain annual cost. Do not add upfront_capex to annualized_capex."},
        {"topic": "Design", "meaning": "Selected equipment, energy requirements and financial assumptions at the reference flow. Resistance is routing metadata only."},
        {"topic": "Costed input", "meaning": "A separate copy retains the input format and all rows. Thirteen pipeline_* columns contain this run's costs, status and error message, including four EUR 2024 columns; unprocessed rows are blank in those columns."},
        {"topic": "Failed routes", "meaning": "Summary includes every attempted route. cost_status=failed records the calculation error in cost_error, with blank costs. Failed routes have no coefficient, component or design rows. Blank costs are not zero-cost connections."},
        {"topic": "Reference costs", "meaning": "Summary cost totals are direct reference calculations. A flow-range fit can differ from them. Reference totals are blank if no direct reference was evaluated."},
        {"topic": "Currency", "meaning": "Existing *_eur* columns and cost_components.value retain EUR 2021. New *_eur2024* columns show the same costs in EUR 2024. Do not add the two price-year representations together. See settings for the resolved factor, source and coverage."},
        {"topic": "Conversion method", "meaning": "Default: AdOpT-NET0 PPI mean(2024)/mean(2021), approximately 1.2089189251. The bundled snapshot has Jan-Dec 2021 and Jan-Nov 2024 only. A documented custom factor can be configured. This is price-year escalation, not discounting or an exchange-rate conversion."},
        {"topic": "Electricity", "meaning": "The configured electricity price is assumed EUR 2021/MWh (default 60 is a scenario assumption). Its EUR 2024 equivalent uses the same PPI factor, not an observed 2024 electricity tariff."},
        {"topic": "Refresh", "meaning": "These are exported Python-model results. Rerun after changing assumptions. EUR 2024 export conversion does not re-size equipment or change the engineering search."},
    ]


def write_results(config, results) -> dict[str, Path]:
    tables = build_tables(config, results)
    output_dir = config.output_dir.resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    names = []
    if config.write_xlsx:
        names.append(config.output_workbook_name)
    input_name = costed_input_name(config.source_path) if config.write_costed_input and config.source_path else None
    if input_name:
        names.append(input_name)
    if config.write_csv:
        names += [f"{name}.csv" for name in tables]
    if config.write_json:
        names.append("pipeline_results.json")
    if not names:
        raise ValueError("No output selected. A costed input copy requires an input file.")
    if len(names) != len(set(names)):
        raise ValueError("Report and costed-input filenames must differ")
    if config.source_path is not None and config.source_path.resolve() in [(output_dir / name) for name in names]:
        raise ValueError("Output would overwrite the input file")
    with tempfile.TemporaryDirectory(prefix=".pipeline-costs-", dir=output_dir) as staging:
        temporary = Path(staging)
        if config.write_xlsx:
            workbook = Workbook()
            workbook.remove(workbook.active)
            for name, rows in tables.items():
                write_sheet(workbook.create_sheet(name), rows)
            write_sheet(workbook.create_sheet("read_me"), report_notes())
            workbook.save(temporary / config.output_workbook_name)
            workbook.close()
        if input_name:
            write_costed_input(temporary / input_name, config, results)
        if config.write_csv:
            for name, rows in tables.items():
                write_csv(temporary / f"{name}.csv", rows)
        if config.write_json:
            payload = {"schema_version": VERSION, "created_utc": datetime.now(timezone.utc).isoformat(),
                       "configuration": asdict(config), "tables": tables}
            (temporary / "pipeline_results.json").write_text(json.dumps(payload, indent=2, default=str, allow_nan=False) + "\n", encoding="utf-8")
        for name in names:
            os.replace(temporary / name, output_dir / name)
    return {name: output_dir / name for name in names}


def write_csv(path, rows):
    keys = list(dict.fromkeys(key for row in rows for key in row))
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=keys)
        writer.writeheader()
        writer.writerows(rows)


def write_sheet(sheet, rows):
    keys = list(dict.fromkeys(key for row in rows for key in row))
    sheet.append(keys)
    for row in rows:
        sheet.append([row.get(key) for key in keys])
    style_sheet(sheet)


def style_sheet(sheet):
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    sheet.row_dimensions[1].height = 72
    for cell in sheet[1]:
        label = str(cell.value)
        cell.font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
        cell.fill = PatternFill("solid", fgColor="214E68")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        width = min(38, max(18, len(label) * .62))
        if label in {"meaning", "value"} and sheet.title in {"settings", "read_me"}:
            width = 100
        elif label in {"metric", "field", "parameter"}:
            width = 48
        elif label == "cost_error":
            width = 100
        sheet.column_dimensions[get_column_letter(cell.column)].width = width
    role_column = next((cell.column for cell in sheet[1] if cell.value == "row_role"), None)
    status_column = next((cell.column for cell in sheet[1] if cell.value == "cost_status"), None)
    for row in sheet.iter_rows(min_row=2):
        sheet.row_dimensions[row[0].row].height = 48 if sheet.title in {"settings", "read_me"} else 32
        if status_column is not None and row[status_column - 1].value == "failed":
            sheet.row_dimensions[row[0].row].height = 96
        total = role_column is not None and row[role_column-1].value == "optimizer_total"
        for cell in row:
            cell.font = Font(name="Calibri", size=11, bold=total)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if total:
                cell.fill = PatternFill("solid", fgColor="E3F0EF")
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.data_type = "s"
            if isinstance(cell.value, (float, int)) and not isinstance(cell.value, bool):
                header = str(sheet.cell(1, cell.column).value)
                cell.number_format = "#,##0.00" if header.endswith(("eur_per_year", "_eur", "eur2024_per_year", "_eur2024")) else "#,##0.000000"
                if header in {"input_data_row", "number_of_booster_stations", "cost_reference_year", "fit_sample_count", "validation_sample_count"}:
                    cell.number_format = "0"
