"""One audit workbook and a costed copy of the original input."""
from __future__ import annotations

from dataclasses import asdict
import json
import math
import os
from pathlib import Path
from uuid import uuid4
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from .costed_input import write_costed_input
from .workbook import _write_gamma2_matrix
from .currency import conversion_metadata, monetary_values_2024, to_eur2024


def _safe(value):
    if isinstance(value, float) and not math.isfinite(value):
        return str(value)
    if isinstance(value, Path):
        return str(value)
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, default=str)
    return value


def _sheet(wb, name, headers, rows):
    if name in wb.sheetnames:
        del wb[name]
    ws = wb.create_sheet(name)
    ws.append(headers)
    for row in rows:
        ws.append([_safe(v) for v in row])
    for cells in ws.iter_rows():
        for cell in cells:
            if isinstance(cell.value, str):
                cell.data_type = "s"  # Identifiers and errors are text, never formulas.
            if isinstance(cell.value, float):
                cell.number_format = "#,##0.000000"
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 45
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col, header in enumerate(headers, 1):
        width = 65 if header in {"cost_error", "cost_note", "value", "meaning"} else min(36, max(16, len(header) + 2))
        ws.column_dimensions[get_column_letter(col)].width = width
    return ws


def summary_rows(results, config):
    for r in results:
        route, c = r.route, r.coefficient
        yield [route.mode, c.from_id if c else route.from_id, c.to_id if c else route.to_id,
               route.source_row, route.source_sheet, route.source_cell,
               c.distance_km if c else route.distance_km, route.annual_flow_t_per_year,
               r.reference_flow_t_per_year, 0.0 if c else None,
               c.gamma2_eur_per_t if c else None, r.annual_cost_eur_per_year, r.average_cost_eur_per_t,
               c.unit_cost_eur_per_t_km if c else None, r.cost_status, r.cost_error, r.cost_note
               ] + monetary_values_2024(r, config.conversion_factor)


SUMMARY_HEADERS = ["mode", "from_id", "to_id", "input_data_row", "source_sheet", "source_cell",
    "distance_km", "input_annual_flow_t_per_year", "reference_flow_t_per_year",
    "fixed_cost_eur_per_year", "flow_cost_eur_per_t", "annual_cost_eur_per_year",
    "average_cost_eur_per_t", "unit_cost_eur_per_t_km", "cost_status", "cost_error", "cost_note",
    "fixed_cost_eur2024_per_year", "flow_cost_eur2024_per_t", "annual_cost_eur2024_per_year",
    "average_cost_eur2024_per_t", "unit_cost_eur2024_per_t_km"]


def _report_tables(wb, config, modes, data, results, *, matrix=False):
    _sheet(wb, "container_cost_coefficients" if matrix else "summary", SUMMARY_HEADERS, summary_rows(results, config))
    factor = config.conversion_factor
    if not matrix:
        coefficients, components = [], []
        for r in results:
            c = r.coefficient
            if c is None:
                continue
            prefix = [r.route.mode, c.from_id, c.to_id, r.route.source_row, r.route.source_cell]
            coefficients.append(prefix + ["distance_regression", 0.0, c.gamma2_eur_per_t, r.reference_flow_t_per_year,
                                          0.0, to_eur2024(c.gamma2_eur_per_t, factor)])
            params = config.parameters(r.route.mode)
            for component, value in (("distance_independent_per_t", params.fixed_eur_per_t),
                                     ("distance_dependent_per_t", params.distance_rate_eur_per_t_km * c.distance_km)):
                annual = None if r.reference_flow_t_per_year is None else value * r.reference_flow_t_per_year
                components.append(prefix + [component, value, annual,
                                              to_eur2024(value, factor), to_eur2024(annual, factor)])
        ids = ["mode", "from_id", "to_id", "input_data_row", "source_cell"]
        _sheet(wb, "coefficients", ids + ["coefficient_method", "fixed_cost_eur_per_year",
               "flow_cost_eur_per_t", "reference_flow_t_per_year", "fixed_cost_eur2024_per_year",
               "flow_cost_eur2024_per_t"], coefficients)
        _sheet(wb, "calculation_components", ids + ["component", "cost_eur_per_t", "annual_cost_eur_per_year",
               "cost_eur2024_per_t", "annual_cost_eur2024_per_year"], components)
    entries = list({**asdict(config), **conversion_metadata(config)}.items()) + [
        ("package_version", "0.3.0"), ("selected_modes", modes), ("resolved_input_layout", data.layout),
        ("cost_equation", "C = F*y + v*Q; F = 0; v = a + b*d"),
        ("successful_routes", sum(r.cost_status == "ok" for r in results)),
        ("failed_routes", sum(r.cost_status == "failed" for r in results)),
        ("routes_with_annual_cost", sum(r.annual_cost_eur_per_year is not None for r in results)),
        ("price_year_note", "Existing *_eur_* fields are EUR 2021. Additional *_eur2024_* fields equal EUR 2021 values times the recorded factor. See conversion_basis for index coverage."),
    ]
    _sheet(wb, "container_cost_model_info" if matrix else "settings", ["parameter", "value"], entries)
    if matrix:
        return
    _sheet(wb, "units", ["field", "unit", "meaning"], [
        ("distance_km", "km", "Physical one-way routed distance; distance is embedded in v exactly once"),
        ("unit_cost_eur_per_t_km", "EUR2021/(tCO2*km)", "UC = a/d + b"),
        ("reference_flow_t_per_year", "tCO2/year", "Transported annual mass; no capture fraction or hourly conversion applied"),
        ("fixed_cost_eur_per_year", "EUR2021/year", "F = 0: no separately modelled annual fixed charge"),
        ("flow_cost_eur_per_t", "EUR2021/tCO2", "v = a + b*d; not an AdOpT CAPEX coefficient per t/h"),
        ("annual_cost_eur_per_year", "EUR2021/year", "v*Q; blank when annual flow is unavailable or calculation failed"),
        ("average_cost_eur_per_t", "EUR2021/tCO2", "C/Q for positive Q; blank at zero or missing Q"),
        ("fixed_cost_eur2024_per_year", "EUR2024/year", "F remains zero; failed routes remain blank"),
        ("flow_cost_eur2024_per_t", "EUR2024/tCO2", "v in EUR 2021 multiplied by the recorded conversion factor"),
        ("annual_cost_eur2024_per_year", "EUR2024/year", "Converted v*Q; missing annual flow remains blank"),
        ("average_cost_eur2024_per_t", "EUR2024/tCO2", "Converted C/Q; blank at zero or missing Q"),
        ("unit_cost_eur2024_per_t_km", "EUR2024/(tCO2*km)", "Converted UC; same routed distance"),
        ("eur2021_to_eur2024_factor", "dimensionless", "Price index ratio or documented custom factor; no exchange rate, discounting or time conversion"),
    ])
    _sheet(wb, "read_me", ["topic", "meaning"], [
        ("Equations", "Truck: UC=5.58/d+0.15; railway: UC=28.9/d+0.07 by default. v=UC*d. Settings record any overrides."),
        ("Fixed term", "5.58 and 28.9 are per-tonne distance-independent charges, not annual fixed costs."),
        ("Annual flow", "A row flow overrides the explicit YAML fallback. Node emissions are never used implicitly."),
        ("Missing flow", "Coefficients remain valid; annual totals and average costs are blank. Enable require_annual_flow to fail these rows."),
        ("Components", "Add the two calculation_components rows per route. They are regression terms, not separate CAPEX/OPEX estimates."),
        ("No double counting", "Use F and v once in annual objective C=F*y+v*Q; do not annualise or multiply v by distance again."),
        ("Availability", "Only successful, available routes may carry flow. A zero F does not enable an unavailable route."),
        ("Boundary", "No extra fleet, terminal, return-trip, loading, access-leg or conditioning charge is added. Confirm regression scope before additions or comparisons."),
        ("Interpretation", "No explicit CAPEX/OPEX split, fleet capacity, economies of scale or off-design engineering calculation is inferred."),
        ("Failures", "Failed costs are blank; inspect cost_status and cost_error. Current result fields are cleared for unprocessed rows."),
        ("Price year", "Published regressions use EUR 2021. Existing EUR fields retain this basis; the added EUR2024 fields contain converted values. Use conversion_basis and settings to audit the factor and month coverage."),
        ("Recalculation", "Outputs contain Python results, not editable Excel formulas; rerun after changing inputs."),
    ])


def write_report(path, config, modes, data, results):
    wb = Workbook()
    wb.remove(wb.active)
    try:
        _report_tables(wb, config, modes, data, results)
        wb.save(path)
    finally:
        wb.close()


def write_matrix_copy(path, config, modes, data, results):
    wb = load_workbook(config.source_path)
    try:
        # Clear generated matrices from previous runs, including unselected modes.
        generated = ("gamma2", "fixed_cost", "annual_cost", "reference_flow", "cost_status",
                     "gamma2_eur2024", "fixed_cost_eur2024", "annual_cost_eur2024")
        for mode in ("truck", "railway"):
            for suffix in generated:
                name = f"{mode}_{suffix}"
                if name in wb.sheetnames:
                    del wb[name]
        for mode in modes:
            ids = data.node_ids_by_mode[mode]
            selected = [r for r in results if r.route.mode == mode]
            _write_gamma2_matrix(wb, mode, ids, [r.coefficient for r in selected if r.coefficient])
            matrices = {}
            for suffix in generated[1:]:
                ws = wb.create_sheet(f"{mode}_{suffix}")
                ws.append(["node_id"] + ids)
                for node in ids:
                    fill = "unavailable" if suffix == "cost_status" else 0.0 if suffix == "gamma2_eur2024" else None
                    ws.append([node] + [fill] * len(ids))
                ws.freeze_panes = "B2"
                matrices[suffix] = ws
            for result in selected:
                cell = result.route.source_cell
                matrices["cost_status"][cell] = result.cost_status
                if result.coefficient:
                    matrices["fixed_cost"][cell] = 0.0
                    matrices["annual_cost"][cell] = result.annual_cost_eur_per_year
                    matrices["reference_flow"][cell] = result.reference_flow_t_per_year
                    converted = monetary_values_2024(result, config.conversion_factor)
                    matrices["fixed_cost_eur2024"][cell] = converted[0]
                    matrices["gamma2_eur2024"][cell] = converted[1]
                    matrices["annual_cost_eur2024"][cell] = converted[2]
                else:
                    wb[f"{mode}_gamma2"][cell] = None
                    matrices["gamma2_eur2024"][cell] = None
        _report_tables(wb, config, modes, data, results, matrix=True)
        wb.save(path)
    finally:
        wb.close()


def write_results(config, modes, data, results):
    directory = Path(config.output_dir)
    directory.mkdir(parents=True, exist_ok=True)
    outputs = {}
    if config.write_xlsx:
        outputs["report"] = directory / config.report_name(modes)
    if config.write_costed_input:
        source = config.source_path
        mode_suffix = f"_{modes[0]}" if len(modes) == 1 else ""
        name = config.output_name(modes) if data.layout == "matrix" else f"{source.stem}{mode_suffix}_costed{source.suffix}"
        outputs["costed_input"] = directory / name
    paths = [p.resolve() for p in outputs.values()]
    if len(set(paths)) != len(paths) or config.source_path.resolve() in paths:
        raise ValueError("Report and costed input must have distinct paths and must not overwrite the input")
    temporary = {}
    try:
        # Prepare both outputs completely before replacing either destination.
        for role, path in outputs.items():
            temp = path.with_name(f".{path.stem}.{uuid4().hex}{path.suffix}")
            temporary[role] = temp
            if role == "report":
                write_report(temp, config, modes, data, results)
            elif data.layout == "matrix":
                write_matrix_copy(temp, config, modes, data, results)
            else:
                write_costed_input(temp, config, results)
        for role, path in outputs.items():
            os.replace(temporary[role], path)
    finally:
        for temp in temporary.values():
            temp.unlink(missing_ok=True)
    return outputs
