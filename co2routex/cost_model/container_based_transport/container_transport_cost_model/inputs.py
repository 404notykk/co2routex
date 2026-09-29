"""Read route tables or directed distance matrices without guessing annual flows."""
from __future__ import annotations

import csv
import math
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from .config import normalize_mode
from .models import InputData, RouteInput


def missing(value) -> bool:
    return value is None or (isinstance(value, str) and not value.strip()) or (
        isinstance(value, float) and math.isnan(value)
    )


def clean_id(value) -> str:
    if missing(value):
        raise ValueError("Node IDs cannot be blank")
    text = str(value).strip()
    return text[:-2] if text.endswith(".0") and text[:-2].isdigit() else text


def _check_headers(headers, required=()):
    if any(not isinstance(h, str) or not h.strip() for h in headers):
        raise ValueError("Input columns must have nonblank text headers")
    if len(headers) != len(set(headers)):
        raise ValueError("Input contains duplicate column names")
    absent = sorted(set(required) - set(headers))
    if absent:
        raise ValueError(f"Missing input columns: {', '.join(absent)}")


def read_routes(config, modes) -> InputData:
    suffix = config.source_path.suffix.lower()
    if suffix == ".xlsx":
        wb = load_workbook(config.source_path, read_only=True, data_only=True)
        try:
            layout = config.input_layout
            if layout == "auto":
                layout = "routes" if config.route_metrics_sheet in wb.sheetnames else (
                    "matrix" if any(config.source_sheet(m) in wb.sheetnames for m in modes) else "routes"
                )
            if layout == "matrix":
                return _read_matrices(wb, config, modes)
            if config.route_metrics_sheet not in wb.sheetnames:
                raise ValueError(f"Workbook has no sheet {config.route_metrics_sheet!r}; set route_metrics_sheet or input_layout: matrix")
            ws = wb[config.route_metrics_sheet]
            rows = ws.iter_rows(values_only=True)
            headers = list(next(rows, ()))
            records = ((i, dict(zip(headers, row))) for i, row in enumerate(rows, 1))
            routes = _table_routes(headers, records, config, modes, ws.title, ws.max_row - 1)
            return InputData(routes, "routes", ws.title)
        finally:
            wb.close()
    if config.input_layout == "matrix":
        raise ValueError("Matrix input requires an XLSX workbook")
    if suffix == ".csv":
        with config.source_path.open(encoding="utf-8-sig", newline="") as stream:
            reader = csv.DictReader(stream)
            headers = list(reader.fieldnames or [])
            records = list(enumerate(reader, 1))
            return InputData(_table_routes(headers, records, config, modes, None, len(records)), "routes")
    if suffix in {".parquet", ".geoparquet"}:
        try:
            import pyarrow.parquet as pq
        except ImportError as exc:
            raise ImportError("Parquet input requires pyarrow; install requirements.txt") from exc
        with pq.ParquetFile(config.source_path) as source:
            headers = source.schema_arrow.names
            relevant = {config.from_id_column, config.to_id_column, config.distance_column,
                        config.annual_flow_column, config.mode_column, config.status_column}
            columns = [h for h in headers if h in relevant]
            total = source.metadata.num_rows
            if config.input_row is not None:
                if config.input_row > total:
                    raise ValueError(f"input_row {config.input_row} exceeds {total} stored rows")
                offset = 0
                for group in range(source.num_row_groups):
                    count = source.metadata.row_group(group).num_rows
                    if config.input_row <= offset + count:
                        table = source.read_row_group(group, columns=columns).select(columns)
                        row = table.slice(config.input_row - 1 - offset, 1).to_pylist()[0]
                        records = [(config.input_row, row)]
                        break
                    offset += count
            else:
                def records_from_batches():
                    index = 0
                    for batch in source.iter_batches(columns=columns, batch_size=65536):
                        for row in batch.select(columns).to_pylist():
                            index += 1
                            yield index, row
                records = records_from_batches()
            return InputData(_table_routes(headers, records, config, modes, None, total), "routes")
    raise ValueError(f"Unsupported input: {suffix}")


def _table_routes(headers, records, config, modes, sheet, total):
    _check_headers(headers, (config.from_id_column, config.to_id_column, config.distance_column))
    if config.input_row is not None and config.input_row > total:
        raise ValueError(f"input_row {config.input_row} exceeds {total} data rows")
    fallback = config.default_transport_mode or (modes[0] if len(modes) == 1 else None)
    if config.mode_column not in headers and fallback is None:
        raise ValueError("No mode column: select one mode with --transport-mode or set default_transport_mode")
    routes = []
    accepted = {str(s).strip().lower() for s in config.accepted_statuses}
    for number, row in records:
        if config.input_row is not None and number != config.input_row:
            continue
        if config.status_column in headers and str(row.get(config.status_column)).strip().lower() not in accepted:
            continue
        value = row.get(config.mode_column) if config.mode_column is not None else None
        if missing(value):
            if fallback is None:
                raise ValueError(f"Data row {number}: transport mode is blank; specify its mode")
            value = fallback
        try:
            mode = normalize_mode(value)
        except ValueError:
            if str(value).strip().lower() in {"pipeline", "ship", "shipping", "barge"}:
                continue
            raise ValueError(f"Data row {number}: unknown mode {value!r}") from None
        if mode not in modes:
            continue
        routes.append(RouteInput(mode, row.get(config.from_id_column), row.get(config.to_id_column),
                                 row.get(config.distance_column), row.get(config.annual_flow_column),
                                 number, sheet))
    if not routes:
        raise ValueError("No eligible truck/railway routes; check mode/status filters and input_row (no substitute row is selected)")
    return routes


def _matrix_rows(ws):
    rows = list(ws.iter_rows(values_only=True))
    ids = [clean_id(v) for v in rows[0][1:]] if rows else []
    row_ids = [clean_id(row[0]) for row in rows[1:]]
    if not ids or ids != row_ids or len(ids) != len(set(ids)):
        raise ValueError(f"{ws.title!r} must be square with unique, matching row/column node IDs")
    return ids, rows


def _read_matrices(wb, config, modes):
    if config.input_row is not None:
        raise ValueError("--row/input_row applies to route tables; matrix rows contain multiple connections")
    reserved = {f"{mode}_{suffix}" for mode in ("truck", "railway")
                for suffix in ("gamma2", "fixed_cost", "annual_cost", "reference_flow", "cost_status",
                               "gamma2_eur2024", "fixed_cost_eur2024", "annual_cost_eur2024")}
    reserved.update({"container_cost_coefficients", "container_cost_model_info"})
    source_names = {config.source_sheet(mode) for mode in modes}
    source_names.update(getattr(config, f"{mode}_annual_flow_sheet") for mode in modes)
    if reserved.intersection(source_names):
        raise ValueError("A source matrix sheet name conflicts with a reserved output sheet name; rename that input sheet")
    routes, nodes = [], {}
    for mode in modes:
        name = config.source_sheet(mode)
        if name not in wb.sheetnames:
            raise ValueError(f"Workbook does not contain sheet {name!r}")
        ids, rows = _matrix_rows(wb[name])
        nodes[mode] = ids
        flow_name = getattr(config, f"{mode}_annual_flow_sheet")
        flows = None
        if flow_name:
            if flow_name not in wb.sheetnames:
                raise ValueError(f"Workbook does not contain annual-flow sheet {flow_name!r}")
            flow_ids, flows = _matrix_rows(wb[flow_name])
            if flow_ids != ids:
                raise ValueError(f"{flow_name!r} must have the same node order as {name!r}")
        for ri, origin in enumerate(ids, 1):
            for ci, destination in enumerate(ids, 1):
                value = rows[ri][ci]
                if missing(value):
                    continue
                try:
                    if not isinstance(value, bool) and float(value) == 0:
                        continue
                except (ValueError, TypeError):
                    pass  # Invalid values become failed route records, not zero costs.
                cell = f"{get_column_letter(ci + 1)}{ri + 1}"
                routes.append(RouteInput(mode, origin, destination, value,
                    flows[ri][ci] if flows is not None else None, ri, name, cell))
    return InputData(routes, "matrix", node_ids_by_mode=nodes)
