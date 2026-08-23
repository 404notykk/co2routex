"""Read distance matrices and write coefficient matrices."""

from __future__ import annotations

import math
import os
import shutil
from pathlib import Path
from uuid import uuid4

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

from .config import TransportCostConfig
from .cost_model import calculate_pair_coefficient
from .models import PairCoefficient


DETAIL_SHEET = "container_cost_coefficients"
INFO_SHEET = "container_cost_model_info"
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def read_coefficients(
    config: TransportCostConfig,
    mode: str,
) -> tuple[list[str], list[PairCoefficient]]:
    workbook = load_workbook(config.workbook_path, read_only=True, data_only=True)
    try:
        sheet_name = config.source_sheet(mode)
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Workbook does not contain sheet {sheet_name!r}")
        sheet = workbook[sheet_name]
        node_ids = [_clean_id(sheet.cell(1, col).value) for col in range(2, sheet.max_column + 1)]
        if not node_ids or any(not node_id for node_id in node_ids):
            raise ValueError(f"{sheet_name!r} must contain node IDs across row 1")
        if len(node_ids) != len(set(node_ids)):
            raise ValueError(f"{sheet_name!r} contains duplicate column node IDs")

        row_ids = [_clean_id(sheet.cell(row, 1).value) for row in range(2, sheet.max_row + 1)]
        if row_ids != node_ids:
            raise ValueError(
                f"{sheet_name!r} must be a square matrix with identical row and column node IDs"
            )

        parameters = config.parameters(mode)
        results: list[PairCoefficient] = []
        for row, from_id in enumerate(row_ids, start=2):
            for col, to_id in enumerate(node_ids, start=2):
                value = sheet.cell(row, col).value
                if value is None or value == "":
                    continue
                try:
                    distance = float(value)
                except (TypeError, ValueError) as exc:
                    raise ValueError(
                        f"Invalid distance in {sheet_name}!{sheet.cell(row, col).coordinate}"
                    ) from exc
                if not math.isfinite(distance) or distance < 0:
                    raise ValueError(
                        f"Distance must be finite and non-negative in "
                        f"{sheet_name}!{sheet.cell(row, col).coordinate}"
                    )
                if distance > 0:
                    results.append(
                        calculate_pair_coefficient(mode, from_id, to_id, distance, parameters)
                    )
        return node_ids, results
    finally:
        workbook.close()


def write_output(
    config: TransportCostConfig,
    modes: tuple[str, ...],
    node_ids_by_mode: dict[str, list[str]],
    coefficients: list[PairCoefficient],
) -> Path:
    config.output_dir.mkdir(parents=True, exist_ok=True)
    output_path = config.output_dir / config.output_name(modes)
    if output_path.resolve() == config.workbook_path.resolve():
        raise ValueError("Output workbook must not overwrite the input workbook")
    temporary_path = output_path.with_name(
        f".{output_path.stem}.{uuid4().hex}{output_path.suffix}"
    )
    try:
        shutil.copy2(config.workbook_path, temporary_path)
        workbook = load_workbook(temporary_path)
        try:
            for mode in modes:
                _write_gamma2_matrix(
                    workbook,
                    mode,
                    node_ids_by_mode[mode],
                    [item for item in coefficients if item.mode == mode],
                )
            _write_details(workbook, coefficients)
            _write_info(workbook, config, modes)
            workbook.save(temporary_path)
        finally:
            workbook.close()
        os.replace(temporary_path, output_path)
    finally:
        temporary_path.unlink(missing_ok=True)
    return output_path


def _write_gamma2_matrix(workbook, mode: str, node_ids: list[str], rows: list[PairCoefficient]) -> None:
    sheet_name = f"{mode}_gamma2"
    if sheet_name in workbook.sheetnames:
        del workbook[sheet_name]
    sheet = workbook.create_sheet(sheet_name)
    sheet.cell(1, 1, "node_id")
    for col, node_id in enumerate(node_ids, start=2):
        sheet.cell(1, col, node_id)
    for row, node_id in enumerate(node_ids, start=2):
        sheet.cell(row, 1, node_id)
        for col in range(2, len(node_ids) + 2):
            sheet.cell(row, col, 0.0)

    positions = {node_id: index + 2 for index, node_id in enumerate(node_ids)}
    for result in rows:
        cell = sheet.cell(positions[result.from_id], positions[result.to_id])
        cell.value = result.gamma2_eur_per_t
        cell.number_format = "0.000000"

    _style_header(sheet, len(node_ids) + 1)
    sheet.freeze_panes = "B2"
    sheet.column_dimensions["A"].width = 14


def _write_details(workbook, coefficients: list[PairCoefficient]) -> None:
    if DETAIL_SHEET in workbook.sheetnames:
        del workbook[DETAIL_SHEET]
    sheet = workbook.create_sheet(DETAIL_SHEET)
    headers = [
        "mode",
        "from_id",
        "to_id",
        "distance_km",
        "unit_cost_eur_per_t_km",
        "gamma1_eur",
        "gamma2_eur_per_t",
        "gamma3_eur_per_km",
        "gamma4_eur_per_t_km",
    ]
    sheet.append(headers)
    for item in coefficients:
        sheet.append([
            item.mode,
            item.from_id,
            item.to_id,
            item.distance_km,
            item.unit_cost_eur_per_t_km,
            item.gamma1_eur,
            item.gamma2_eur_per_t,
            item.gamma3_eur_per_km,
            item.gamma4_eur_per_t_km,
        ])
    _style_header(sheet, len(headers))
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    widths = [12, 14, 14, 15, 25, 15, 20, 20, 22]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=2, min_col=4, max_col=9):
        for cell in row:
            cell.number_format = "0.000000"


def _write_info(workbook, config: TransportCostConfig, modes: tuple[str, ...]) -> None:
    if INFO_SHEET in workbook.sheetnames:
        del workbook[INFO_SHEET]
    sheet = workbook.create_sheet(INFO_SHEET)
    sheet.append(["parameter", "value"])
    entries = [
        ("modes", ", ".join(modes)),
        ("cost_equation", "cost = gamma2 * transported_CO2"),
        ("gamma1", 0),
        ("gamma3", 0),
        ("gamma4", 0),
        (
            "truck_UC",
            f"{config.truck_fixed_eur_per_t} / distance_km + "
            f"{config.truck_distance_rate_eur_per_t_km} EUR/(t*km)",
        ),
        (
            "truck_gamma2",
            f"{config.truck_fixed_eur_per_t} + "
            f"{config.truck_distance_rate_eur_per_t_km} * distance_km EUR/t",
        ),
        (
            "railway_UC",
            f"{config.railway_fixed_eur_per_t} / distance_km + "
            f"{config.railway_distance_rate_eur_per_t_km} EUR/(t*km)",
        ),
        (
            "railway_gamma2",
            f"{config.railway_fixed_eur_per_t} + "
            f"{config.railway_distance_rate_eur_per_t_km} * distance_km EUR/t",
        ),
        ("currency", config.currency),
        ("cost_reference_year", config.cost_reference_year),
        ("source_note", config.source_note),
        ("time_basis", "Multiply EUR/t gamma2 by transported tonnes in the chosen period"),
        ("zero_semantics", "0 = unavailable connection"),
    ]
    for entry in entries:
        sheet.append(entry)
    _style_header(sheet, 2)
    sheet.column_dimensions["A"].width = 24
    sheet.column_dimensions["B"].width = 78
    sheet.freeze_panes = "A2"


def _style_header(sheet, columns: int) -> None:
    for col in range(1, columns + 1):
        cell = sheet.cell(1, col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT


def _clean_id(value) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if text.endswith(".0") and text[:-2].isdigit():
        return text[:-2]
    return text


def _column_letter(index: int) -> str:
    letters = ""
    while index:
        index, remainder = divmod(index - 1, 26)
        letters = chr(65 + remainder) + letters
    return letters
