"""Preserve all input columns, rows, geometry and metadata while appending costs."""
from __future__ import annotations
from bisect import bisect_left
import csv
import json
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.table import TableColumn
from .currency import conversion_metadata, monetary_values_2024

COSTED_INPUT_COLUMNS = (
    "container_transport_mode", "container_coefficient_method",
    "container_reference_flow_t_per_year", "container_fixed_cost_eur_per_year",
    "container_flow_cost_eur_per_t", "container_annual_cost_eur_per_year",
    "container_average_cost_eur_per_t", "container_cost_status",
    "container_cost_error", "container_cost_note",
    "container_fixed_cost_eur2024_per_year", "container_flow_cost_eur2024_per_t",
    "container_annual_cost_eur2024_per_year", "container_average_cost_eur2024_per_t",
    "container_cost_reference_year", "container_converted_price_year",
    "container_eur2021_to_eur2024_factor", "container_currency_conversion_basis",
)
TEXT_COLUMNS = frozenset(("container_transport_mode", "container_coefficient_method",
                         "container_cost_status", "container_cost_error", "container_cost_note",
                         "container_currency_conversion_basis"))
INTEGER_COLUMNS = frozenset(("container_cost_reference_year", "container_converted_price_year"))


def result_columns(result, config):
    values = dict.fromkeys(COSTED_INPUT_COLUMNS)
    values.update(container_transport_mode=result.route.mode, container_cost_status=result.cost_status,
                  container_cost_error=result.cost_error, container_cost_note=result.cost_note)
    metadata = conversion_metadata(config)
    values.update(container_cost_reference_year=2021, container_converted_price_year=2024,
                  container_eur2021_to_eur2024_factor=config.conversion_factor,
                  container_currency_conversion_basis=metadata["conversion_basis"])
    if result.coefficient is not None:
        values.update(container_coefficient_method="distance_regression",
                      container_reference_flow_t_per_year=result.reference_flow_t_per_year,
                      container_fixed_cost_eur_per_year=0.0,
                      container_flow_cost_eur_per_t=result.coefficient.gamma2_eur_per_t,
                      container_annual_cost_eur_per_year=result.annual_cost_eur_per_year,
                      container_average_cost_eur_per_t=result.average_cost_eur_per_t)
        converted = monetary_values_2024(result, config.conversion_factor)
        values.update(zip(("container_fixed_cost_eur2024_per_year", "container_flow_cost_eur2024_per_t",
                           "container_annual_cost_eur2024_per_year", "container_average_cost_eur2024_per_t"), converted[:4]))
    return values


def write_costed_input(path, config, results):
    values = {result.route.source_row - 1: result_columns(result, config) for result in results}
    if len(values) != len(results):
        raise ValueError("Duplicate physical source rows")
    extension = config.source_path.suffix.lower()
    if extension in {".parquet", ".geoparquet"}:
        _write_parquet(path, config.source_path, values, conversion_metadata(config))
    elif extension == ".xlsx":
        _write_xlsx(path, config, values)
    elif extension == ".csv":
        _write_csv(path, config.source_path, values)
    else:
        raise ValueError(f"Unsupported costed-input format: {extension}")


def _write_csv(path, source, values):
    with source.open(encoding="utf-8-sig", newline="") as stream, path.open("w", encoding="utf-8-sig", newline="") as output:
        reader = csv.DictReader(stream)
        headers = list(reader.fieldnames or [])
        headers += [key for key in COSTED_INPUT_COLUMNS if key not in headers]
        writer = csv.DictWriter(output, fieldnames=headers)
        writer.writeheader()
        count = 0
        for index, row in enumerate(reader):
            row.update(dict.fromkeys(COSTED_INPUT_COLUMNS))
            row.update(values.get(index, {}))
            writer.writerow(row)
            count += 1
        if values and max(values) >= count:
            raise ValueError("A calculated row is outside the original CSV input")


def _write_xlsx(path, config, values):
    workbook = load_workbook(config.source_path)
    try:
        sheet = workbook[config.route_metrics_sheet]
        original_width, original_height = sheet.max_column, sheet.max_row
        if values and max(values) + 2 > original_height:
            raise ValueError("A calculated row is outside the original Excel input")
        headers = {str(cell.value).strip(): cell.column for cell in sheet[1] if cell.value is not None}
        for key in COSTED_INPUT_COLUMNS:
            if key not in headers:
                headers[key] = sheet.max_column + 1
                sheet.cell(1, headers[key], key)
            column = headers[key]
            cell = sheet.cell(1, column)
            cell.font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
            cell.fill = PatternFill("solid", fgColor="214E68")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            sheet.column_dimensions[get_column_letter(column)].width = 100 if key == "container_cost_error" else 32
            for row in range(2, original_height + 1):
                cell = sheet.cell(row, column)
                cell.value = values.get(row - 2, {}).get(key)
                if key == "container_cost_error" and cell.value is not None:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                    sheet.row_dimensions[row].height = max(sheet.row_dimensions[row].height or 15, 96)
                if key not in TEXT_COLUMNS and key not in INTEGER_COLUMNS:
                    cell.number_format = "#,##0.000000"
        # Include the appended fields in full-input Excel tables so sorting keeps
        # route identifiers and results together. Unrelated tables are unchanged.
        for table in sheet.tables.values():
            left, top, right, bottom = range_boundaries(table.ref)
            if (left, top, right, bottom) == (1, 1, original_width, original_height):
                for column in range(original_width + 1, sheet.max_column + 1):
                    table.tableColumns.append(TableColumn(id=column, name=str(sheet.cell(1, column).value)))
                table.ref = f"A1:{get_column_letter(sheet.max_column)}{original_height}"
                if table.autoFilter is not None:
                    table.autoFilter.ref = table.ref
        sheet.freeze_panes = sheet.freeze_panes or "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.row_dimensions[1].height = max(sheet.row_dimensions[1].height or 15, 80)
        workbook.save(path)
    finally:
        workbook.close()


def _write_parquet(path, source, values, conversion):
    import pyarrow as pa
    import pyarrow.parquet as pq
    with pq.ParquetFile(source) as original:
        if values and max(values) >= original.metadata.num_rows:
            raise ValueError("A calculated row is outside the original Parquet input")
        schema = original.schema_arrow
        new_fields = {key: pa.field(key, pa.string() if key in TEXT_COLUMNS else
                                   pa.int16() if key in INTEGER_COLUMNS else pa.float64())
                      for key in COSTED_INPUT_COLUMNS}
        metadata = dict(schema.metadata or {})
        metadata[b"co2routex.container_cost_model"] = json.dumps(conversion).encode("utf-8")
        # Preserve GeoParquet metadata verbatim. Extend pandas metadata so saved
        # indexes (including duplicate labels) and original dtypes remain intact.
        if b"pandas" in metadata:
            pandas_meta = json.loads(metadata[b"pandas"])
            if any(key in pandas_meta.get("index_columns", []) for key in COSTED_INPUT_COLUMNS):
                raise ValueError("An input index name conflicts with an output cost column")
            descriptions = {key: {"name": key, "field_name": key,
                "pandas_type": "unicode" if key in TEXT_COLUMNS else "int16" if key in INTEGER_COLUMNS else "float64",
                "numpy_type": "object" if key in TEXT_COLUMNS else "int16" if key in INTEGER_COLUMNS else "float64", "metadata": None}
                for key in COSTED_INPUT_COLUMNS}
            columns = pandas_meta.get("columns", [])
            existing = {entry["field_name"] for entry in columns}
            pandas_meta["columns"] = [descriptions.get(entry["field_name"], entry) for entry in columns]
            pandas_meta["columns"] += [descriptions[key] for key in COSTED_INPUT_COLUMNS if key not in existing]
            metadata[b"pandas"] = json.dumps(pandas_meta).encode("utf-8")
        fields = [new_fields.get(field.name, field) for field in schema]
        fields += [new_fields[key] for key in COSTED_INPUT_COLUMNS if key not in schema.names]
        target_schema = pa.schema(fields, metadata=metadata)
        positions = sorted(values)
        offset = 0
        # Stream full input batches, retaining every original column, geometry,
        # nested field and physical row. Only cost fields of selected rows change.
        with pq.ParquetWriter(path, target_schema, compression="snappy") as writer:
            for batch in original.iter_batches(batch_size=65536):
                indices = positions[bisect_left(positions, offset):bisect_left(positions, offset + batch.num_rows)]
                added = {}
                for key, field in new_fields.items():
                    column = [None] * batch.num_rows
                    for index in indices:
                        column[index - offset] = values[index][key]
                    added[key] = pa.array(column, type=field.type)
                arrays = [added[field.name] if field.name in added
                          else batch.column(batch.schema.get_field_index(field.name)) for field in target_schema]
                writer.write_table(pa.Table.from_arrays(arrays, schema=target_schema))
                offset += batch.num_rows
