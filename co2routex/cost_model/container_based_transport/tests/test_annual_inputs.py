import csv
from dataclasses import replace
import hashlib
import json
from pathlib import Path
import subprocess
import sys

import pandas as pd
import pyarrow as pa
import pyarrow.parquet as pq
import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
from openpyxl.worksheet.table import Table
from openpyxl.utils import get_column_letter

from container_transport_cost_model import TransportCostConfig, RouteInput, calculate_route, run_with_outputs
from container_transport_cost_model.costed_input import COSTED_INPUT_COLUMNS
from container_transport_cost_model.currency import DEFAULT_FACTOR, conversion_metadata


def rows(ws):
    values = list(ws.values)
    return [dict(zip(values[0], row)) for row in values[1:]]


def make_csv(path, records):
    headers = list(dict.fromkeys(key for row in records for key in row))
    with path.open('w', newline='') as stream:
        writer = csv.DictWriter(stream, fieldnames=headers)
        writer.writeheader()
        writer.writerows(records)


def route(mode='truck', distance=100, flow=100000, **extra):
    return dict(from_id='001', to_id='002', mode=mode, distance_km=distance,
                annual_flow_t_per_year=flow, **extra)


def test_source_equations_and_annual_accounting():
    config = TransportCostConfig()
    truck = calculate_route(RouteInput('truck', 'A', 'B', 100, 100000), config)
    rail = calculate_route(RouteInput('railway', 'A', 'B', 100, 100000), config)
    assert truck.coefficient.unit_cost_eur_per_t_km == pytest.approx(0.2058)
    assert truck.coefficient.gamma2_eur_per_t == pytest.approx(20.58)
    assert truck.annual_cost_eur_per_year == pytest.approx(2058000)
    assert rail.annual_cost_eur_per_year == pytest.approx(3590000)
    larger = calculate_route(RouteInput('truck', 'A', 'B', 100, 200000), config)
    assert larger.coefficient == truck.coefficient
    assert larger.annual_cost_eur_per_year == pytest.approx(2 * truck.annual_cost_eur_per_year)


def test_xlsx_preserves_originals_and_records_per_route_failures(tmp_path):
    source = tmp_path / 'routes.xlsx'
    wb = Workbook()
    ws = wb.active
    ws.title = 'links'
    headers = ['from_id', 'to_id', 'mode', 'distance_km', 'annual_flow_t_per_year', 'comment', 'container_flow_cost_eur_per_t']
    ws.append(headers)
    ws.append(['001', '002', 'truck', 100, 100000, '=1+2', 999])
    ws.append(['003', '004', 'train', -1, 100000, 'bad distance', 999])
    ws.append(['005', '006', 'pipeline', 200, 100000, 'excluded', 999])
    ws['A2'].fill = PatternFill('solid', fgColor='FFFF00')
    ws.add_table(Table(displayName='Routes', ref='A1:G4'))
    wb.create_sheet('notes')['A1'] = 'preserve this sheet'
    wb.save(source)
    before = source.read_bytes()
    config = TransportCostConfig(input_path=source, output_dir=tmp_path/'out', route_metrics_sheet='links')
    paths = run_with_outputs(config)
    assert source.read_bytes() == before
    result = load_workbook(paths['costed_input'])
    values = rows(result['links'])
    assert values[0]['from_id'] == '001'
    assert values[0]['comment'] == '=1+2'
    assert values[0]['container_annual_cost_eur_per_year'] == pytest.approx(2058000)
    assert values[1]['container_cost_status'] == 'failed'
    assert values[1]['container_flow_cost_eur_per_t'] is None
    assert values[2]['container_cost_status'] is None
    assert values[2]['container_flow_cost_eur_per_t'] is None
    assert result['links']['A2'].fill.fgColor.rgb == '00FFFF00'
    assert result['notes']['A1'].value == 'preserve this sheet'
    assert result['links'].tables['Routes'].ref == f'A1:{get_column_letter(6 + len(COSTED_INPUT_COLUMNS))}4'
    assert values[0]['container_annual_cost_eur2024_per_year'] == pytest.approx(2487955.1477774656)
    assert values[1]['container_flow_cost_eur2024_per_t'] is None
    assert values[2]['container_eur2021_to_eur2024_factor'] is None
    result.close()
    report = load_workbook(paths['report'], data_only=True)
    assert len(rows(report['summary'])) == 2
    assert len(rows(report['coefficients'])) == 1
    components = rows(report['calculation_components'])
    assert sum(r['annual_cost_eur_per_year'] for r in components) == pytest.approx(2058000)
    assert sum(r['annual_cost_eur2024_per_year'] for r in components) == pytest.approx(2487955.1477774656)
    summary = rows(report['summary'])[0]
    assert summary['unit_cost_eur2024_per_t_km'] == pytest.approx(0.2058 * DEFAULT_FACTOR)
    settings = dict(report['settings'].values)
    assert settings['source_price_year'] == 2021
    assert settings['ppi_2024_months'] == 11
    report.close()


def test_parquet_row_position_metadata_geometry_and_stale_results(tmp_path):
    source = tmp_path / 'routes.geoparquet'
    df = pd.DataFrame([route('pipeline', 40), route('rail', 500), route('truck', 100)], index=[9, 9, 2])
    df.index.name = 'saved_index'
    df['container_flow_cost_eur_per_t'] = ['stale'] * 3
    df['geometry'] = [b'one', b'two', b'three']
    df['nested'] = [[1, 2], [3], []]
    table = pa.Table.from_pandas(df)
    geometry = table.schema.field('geometry').with_metadata({b'example_field_metadata': b'unchanged'})
    fields = [geometry if f.name == 'geometry' else f for f in table.schema]
    metadata = dict(table.schema.metadata)
    metadata.update({b'geo': b'{"version":"1.0.0","primary_column":"geometry","columns":{}}', b'custom': b'keep'})
    table = pa.Table.from_arrays(table.columns, schema=pa.schema(fields, metadata=metadata))
    pq.write_table(table, source, row_group_size=1)
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    config = TransportCostConfig(input_path=source, output_dir=tmp_path/'out', input_row=2)
    output = run_with_outputs(config)['costed_input']
    assert hashlib.sha256(source.read_bytes()).hexdigest() == digest
    result = pq.read_table(output)
    for name in table.column_names:
        if name not in COSTED_INPUT_COLUMNS:
            assert result[name].equals(table[name])
            assert result.schema.field(name) == table.schema.field(name)
    assert result.schema.metadata[b'geo'] == metadata[b'geo']
    assert result.schema.metadata[b'custom'] == b'keep'
    conversion = json.loads(result.schema.metadata[b'co2routex.container_cost_model'])
    assert conversion['eur2021_to_eur2024_factor'] == pytest.approx(1.2089189250619368)
    assert conversion['ppi_2024_months'] == 11
    costed = result.to_pandas()
    assert costed.index.tolist() == [9, 9, 2]
    assert costed.index.name == 'saved_index'
    assert costed['container_cost_status'].tolist() == [None, 'ok', None]
    assert costed.iloc[1]['container_flow_cost_eur_per_t'] == pytest.approx(63.9)
    assert costed.iloc[1]['container_annual_cost_eur_per_year'] == pytest.approx(6390000)
    assert pd.isna(costed.iloc[0]['container_flow_cost_eur_per_t'])
    assert pd.isna(costed.iloc[2]['container_flow_cost_eur_per_t'])
    assert costed.iloc[1]['container_annual_cost_eur2024_per_year'] == pytest.approx(6390000 * DEFAULT_FACTOR)
    assert costed.iloc[1]['container_cost_reference_year'] == 2021
    assert costed.iloc[1]['container_converted_price_year'] == 2024
    assert pd.isna(costed.iloc[0]['container_annual_cost_eur2024_per_year'])
    with pytest.raises(ValueError, match='No eligible'):
        run_with_outputs(replace(config, input_row=1))


def test_csv_missing_zero_negative_flows_and_mode_aliases(tmp_path):
    source = tmp_path/'routes.csv'
    make_csv(source, [route(flow=''), route('railway', flow=0), route(flow=-5), route('train', 100, 10)])
    config = TransportCostConfig(input_path=source, output_dir=tmp_path/'out')
    output = run_with_outputs(config)['costed_input']
    with output.open(encoding='utf-8-sig', newline='') as stream:
        result = list(csv.DictReader(stream))
    assert result[0]['from_id'] == '001'
    assert result[0]['container_cost_status'] == 'ok'
    assert result[0]['container_annual_cost_eur_per_year'] == ''
    assert result[0]['container_annual_cost_eur2024_per_year'] == ''
    assert float(result[0]['container_flow_cost_eur2024_per_t']) == pytest.approx(20.58 * DEFAULT_FACTOR)
    assert 'Coefficients only' in result[0]['container_cost_note']
    assert float(result[1]['container_annual_cost_eur_per_year']) == 0
    assert result[1]['container_average_cost_eur_per_t'] == ''
    assert float(result[1]['container_annual_cost_eur2024_per_year']) == 0
    assert result[1]['container_average_cost_eur2024_per_t'] == ''
    assert result[2]['container_cost_status'] == 'failed'
    assert result[2]['container_flow_cost_eur_per_t'] == ''
    assert result[2]['container_flow_cost_eur2024_per_t'] == ''
    assert float(result[3]['container_annual_cost_eur_per_year']) == pytest.approx(359)
    # Explicit fallback applies to missing Q, not zero or invalid Q.
    output = run_with_outputs(replace(config, annual_flow_t_per_year=100))['costed_input']
    with output.open(encoding='utf-8-sig') as stream:
        result = list(csv.DictReader(stream))
    assert float(result[0]['container_annual_cost_eur_per_year']) == pytest.approx(2058)
    assert float(result[1]['container_annual_cost_eur_per_year']) == 0
    assert result[2]['container_cost_status'] == 'failed'


def test_matrix_flow_mapping_and_failed_unavailable_distinction(tmp_path):
    source = tmp_path/'matrix.xlsx'
    wb = Workbook()
    ws = wb.active
    ws.title = 'truck'
    for row in [['node_id','A','B'],['A',0,100],['B',-1,0]]:
        ws.append(row)
    q = wb.create_sheet('truck_q')
    for row in [['node_id','A','B'],['A',0,100000],['B',200000,0]]:
        q.append(row)
    wb.save(source)
    config = TransportCostConfig(workbook_path=source, truck_annual_flow_sheet='truck_q', output_dir=tmp_path/'out')
    output = run_with_outputs(config, ('truck',))['costed_input']
    wb = load_workbook(output, data_only=True)
    assert wb['truck_annual_cost']['C2'].value == pytest.approx(2058000)
    assert wb['truck_fixed_cost']['C2'].value == 0
    assert wb['truck_cost_status']['B2'].value == 'unavailable'
    assert wb['truck_cost_status']['B3'].value == 'failed'
    assert wb['truck_gamma2']['B2'].value == 0  # Legacy unavailable sentinel.
    assert wb['truck_gamma2']['B3'].value is None
    assert wb['truck_gamma2_eur2024']['C2'].value == pytest.approx(20.58 * DEFAULT_FACTOR)
    assert wb['truck_gamma2_eur2024']['B2'].value == 0
    assert wb['truck_gamma2_eur2024']['B3'].value is None
    assert wb['truck_annual_cost_eur2024']['C2'].value == pytest.approx(2487955.1477774656)
    assert wb['truck_fixed_cost_eur2024']['C2'].value == 0
    wb.close()


def test_all_failed_still_exports_and_single_mode_inference(tmp_path):
    source = tmp_path/'routes.csv'
    make_csv(source, [dict(from_id='A',to_id='B',distance_km=10)])
    config = TransportCostConfig(input_path=source, output_dir=tmp_path/'out', require_annual_flow=True)
    with pytest.raises(ValueError, match='No mode column'):
        run_with_outputs(config)
    paths = run_with_outputs(config, ('truck',))
    wb = load_workbook(paths['report'], data_only=True)
    assert rows(wb['summary'])[0]['cost_status'] == 'failed'
    assert rows(wb['coefficients']) == []
    wb.close()


@pytest.mark.parametrize('value', [-1, float('inf'), float('nan'), True])
def test_invalid_coefficients_rejected(value):
    with pytest.raises(ValueError):
        TransportCostConfig(truck_fixed_eur_per_t=value).validate(require_input=False)


def test_status_filters_and_rerun_clear_unselected_costs(tmp_path):
    source = tmp_path/'routes.parquet'
    pq.write_table(pa.Table.from_pylist([route(status='ok'), route(status='failed'), route('railway', status='ok')]), source)
    cfg = TransportCostConfig(input_path=source, output_dir=tmp_path/'out')
    first = run_with_outputs(cfg)['costed_input']
    second = run_with_outputs(replace(cfg, input_path=first, input_row=3))['costed_input']
    result = pq.read_table(second).to_pylist()
    assert [r['container_cost_status'] for r in result] == [None,None,'ok']
    assert result[0]['container_flow_cost_eur_per_t'] is None
    assert result[0]['container_flow_cost_eur2024_per_t'] is None
    assert result[0]['container_eur2021_to_eur2024_factor'] is None
    assert result[1]['status'] == 'failed'


def test_cli_input_override_and_file_collision_protection(tmp_path):
    source = tmp_path/'input.csv'
    make_csv(source, [route()])
    yaml = tmp_path/'config.yaml'
    yaml.write_text('input_path: nonexistent.csv\n')
    script = Path(__file__).resolve().parents[1]/'run_transport_cost_model.py'
    run = subprocess.run([sys.executable, str(script), '--config', str(yaml), '--input', str(source),
                          '--output-dir', str(tmp_path/'out'), '--row', '1'], capture_output=True, text=True)
    assert run.returncode == 0, run.stderr
    assert (tmp_path/'out'/'input_costed.csv').exists()
    config = TransportCostConfig(input_path=source, output_dir=tmp_path/'out', output_workbook_name='input_costed.xlsx')
    xlsx = tmp_path/'input.xlsx'
    wb=Workbook()
    wb.active.title='transport_route_metrics'
    wb.active.append(list(route().keys()))
    wb.active.append(list(route().values()))
    wb.save(xlsx)
    with pytest.raises(ValueError, match='distinct paths'):
        run_with_outputs(replace(config, input_path=xlsx))


def test_reserved_matrix_sheet_cannot_overwrite_source_flow(tmp_path):
    source = tmp_path/'matrix.xlsx'
    wb = Workbook()
    for ws, name in ((wb.active, 'truck'), (wb.create_sheet(), 'truck_reference_flow')):
        ws.title = name
        for row in [['node_id','A','B'],['A',0,100],['B',0,0]]:
            ws.append(row)
    wb.save(source)
    config = TransportCostConfig(input_path=source, truck_annual_flow_sheet='truck_reference_flow', output_dir=tmp_path/'out')
    with pytest.raises(ValueError, match='reserved output'):
        run_with_outputs(config, ('truck',))


@pytest.mark.parametrize('selected_mode,expected_statuses', [
    ('truck', ['ok', None]), ('railway', [None, 'ok']),
])
def test_routing_mode_aliases_in_mixed_parquet(tmp_path, selected_mode, expected_statuses):
    source = tmp_path/'routed_modes.parquet'
    pq.write_table(pa.Table.from_pylist([route('road', 100, 100000), route('rail', 500, 100000)]), source)
    config = TransportCostConfig(input_path=source, output_dir=tmp_path/'out')
    output = run_with_outputs(config, (selected_mode,))['costed_input']
    values = pq.read_table(output).to_pylist()
    assert [row['mode'] for row in values] == ['road', 'rail']
    assert [row['container_cost_status'] for row in values] == expected_statuses
    index = 0 if selected_mode == 'truck' else 1
    assert values[index]['container_transport_mode'] == selected_mode
    assert values[index]['container_annual_cost_eur_per_year'] == pytest.approx(
        2058000 if selected_mode == 'truck' else 6390000
    )
    assert values[index]['container_annual_cost_eur2024_per_year'] == pytest.approx(
        values[index]['container_annual_cost_eur_per_year'] * DEFAULT_FACTOR
    )
    assert values[1-index]['container_annual_cost_eur2024_per_year'] is None


def test_adopt_snapshot_and_legacy_yaml_migration(tmp_path):
    config = tmp_path/'config.yaml'
    config.write_text('cost_reference_year: null\n')
    loaded = TransportCostConfig.from_yaml(config, require_input=False)
    assert loaded.cost_reference_year == 2021
    assert loaded.conversion_factor == pytest.approx(1.2089189250619368, rel=1e-14)
    meta = conversion_metadata(loaded)
    assert meta['ppi_2021_months'] == 12
    assert meta['ppi_2024_months'] == 11
    assert 'Jan-Nov' in meta['conversion_basis']


def test_custom_factor_is_applied_once_and_its_source_is_exported(tmp_path):
    source = tmp_path/'routes.parquet'
    pq.write_table(pa.Table.from_pylist([route('road'), route('rail')]), source)
    config = TransportCostConfig(input_path=source, output_dir=tmp_path/'out',
                                eur2021_to_eur2024_factor=1.25,
                                eur2024_factor_source='Synthetic test factor; not an observed price index')
    outputs = run_with_outputs(config)
    table = pq.read_table(outputs['costed_input'])
    results = table.to_pylist()
    assert results[0]['container_annual_cost_eur_per_year'] == pytest.approx(2058000)
    assert results[0]['container_annual_cost_eur2024_per_year'] == pytest.approx(2572500)
    assert results[1]['container_annual_cost_eur2024_per_year'] == pytest.approx(4487500)
    assert results[0]['container_currency_conversion_basis'] == config.eur2024_factor_source
    assert 'ppi_2024_months' not in json.loads(table.schema.metadata[b'co2routex.container_cost_model'])
    # A costed-input rerun recomputes from original distance/flow, not the old monetary fields.
    again = run_with_outputs(replace(config, input_path=outputs['costed_input']))
    repeated = pq.read_table(again['costed_input']).to_pylist()
    assert repeated[0]['container_annual_cost_eur2024_per_year'] == pytest.approx(2572500)
    report = load_workbook(outputs['report'], data_only=True)
    assert rows(report['coefficients'])[0]['flow_cost_eur2024_per_t'] == pytest.approx(25.725)
    report.close()


@pytest.mark.parametrize('factor', [0, -1, float('inf'), float('nan'), True])
def test_invalid_conversion_factor_is_rejected(factor):
    with pytest.raises(ValueError):
        TransportCostConfig(eur2021_to_eur2024_factor=factor,
                            eur2024_factor_source='Test').validate(require_input=False)


def test_factor_requires_source_and_wrong_base_year_is_rejected():
    with pytest.raises(ValueError, match='Document a custom factor'):
        TransportCostConfig(eur2021_to_eur2024_factor=1.2).validate(require_input=False)
    with pytest.raises(ValueError, match='must be 2021'):
        TransportCostConfig(cost_reference_year=2024).validate(require_input=False)


def test_conversion_overflow_is_a_route_failure_not_infinite_output(tmp_path):
    source = tmp_path/'routes.csv'
    make_csv(source, [route(flow=1e300)])
    config = TransportCostConfig(input_path=source, output_dir=tmp_path/'out',
                                eur2021_to_eur2024_factor=1e10, eur2024_factor_source='Overflow test')
    outputs = run_with_outputs(config)
    with outputs['costed_input'].open(encoding='utf-8-sig') as stream:
        result = next(csv.DictReader(stream))
    assert result['container_cost_status'] == 'failed'
    assert result['container_annual_cost_eur2024_per_year'] == ''
    assert 'EUR 2024 cost is nonfinite' in result['container_cost_error']
