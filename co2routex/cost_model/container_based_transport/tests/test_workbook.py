from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from container_transport_cost_model.config import TransportCostConfig
from container_transport_cost_model.runner import run


def _make_workbook(path: Path) -> None:
    workbook = Workbook()
    nodes = workbook.active
    nodes.title = "nodes"
    nodes.append(["node_id", "annual_flux"])
    nodes.append(["A", 100000])
    nodes.append(["B", 200000])
    for mode, distance in (("truck", 100), ("railway", 500)):
        sheet = workbook.create_sheet(mode)
        sheet.append(["node_id", "A", "B"])
        sheet.append(["A", 0, distance])
        sheet.append(["B", 0, 0])
    workbook.save(path)


def test_combined_run_preserves_distances_and_adds_gamma2(tmp_path):
    input_path = tmp_path / "node_metrics_routed.xlsx"
    _make_workbook(input_path)
    config = TransportCostConfig(workbook_path=input_path, output_dir=tmp_path / "out")

    output = run(config)

    workbook = load_workbook(output, data_only=True)
    try:
        assert workbook["truck"]["C2"].value == 100
        assert workbook["railway"]["C2"].value == 500
        assert workbook["truck_gamma2"]["C2"].value == pytest.approx(20.58)
        assert workbook["railway_gamma2"]["C2"].value == pytest.approx(63.9)
        assert workbook["truck_gamma2"]["B2"].value == 0
        assert workbook["railway_gamma2"]["B2"].value == 0
        assert "container_cost_coefficients" in workbook.sheetnames
        assert "container_cost_model_info" in workbook.sheetnames
    finally:
        workbook.close()


def test_single_mode_uses_separate_output_name(tmp_path):
    input_path = tmp_path / "node_metrics_routed.xlsx"
    _make_workbook(input_path)
    config = TransportCostConfig(workbook_path=input_path, output_dir=tmp_path / "out")

    output = run(config, ("truck",))

    assert output.name == "node_metrics_truck_costed.xlsx"
    workbook = load_workbook(output, read_only=True)
    try:
        assert "truck_gamma2" in workbook.sheetnames
        assert "railway_gamma2" not in workbook.sheetnames
    finally:
        workbook.close()

